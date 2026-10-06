#!/usr/bin/env python3
"""69-decisions_jf_1006b.py — décisions de JF du 2026-10-06 (après-midi) :
  1. Figure S3 validée (rendu, italique des phylums) : rien à modifier ;
  2. retrait de la mention « the contract that also funded Guivier et al. [31] » validé : note grise mise à jour ;
  3. pas d'analyse d'abondance différentielle (décision 6) : note grise du §8.7 mise à jour ;
  4. confusion taille / catégorie : une phrase dans les limites (Discussion), « Four limits » -> « Five limits ».
     Chiffres : metadata/genotypes_verifies_sept_180.csv (md5 aab2f05c9ac0433bc8210739697e5298), longueur totale
     moyenne Cn 24.4 (n = 59), Hy 19.4 (42), Pt 16.0 (78, une valeur manquante) ; Kruskal-Wallis p = 0.00448.
python-docx requis (absent sur meso) : exécuté hors cluster, fichier déposé avec garde md5.
Usage : 69-decisions_jf_1006b.py <dossier_entree> <dossier_sortie>
"""
import sys, copy, hashlib
import docx
from docx.text.run import Run

IN, OUTD = sys.argv[1], sys.argv[2]
assert hashlib.md5(open(f"{IN}/Article.docx", "rb").read()).hexdigest() == "4398edc2c5d2927b6f419f0b0c97ac06"
A = docx.Document(f"{IN}/Article.docx"); P = A.paragraphs


def find1(sub):
    k = [i for i, p in enumerate(P) if sub in p.text]
    assert len(k) == 1, (sub[:50], k); return P[k[0]]


def replace_once(par, old, new):
    hits = [r for r in par.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (old[:60], len(hits))
    hits[0].text = hits[0].text.replace(old, new)

# 3. abondance différentielle
p52 = find1("« The choice of a differential-abundance method remains to be fixed. »")
replace_once(p52, "Aucune analyse d'abondance différentielle n'est rapportée ; à trancher par JF.]",
             "Aucune analyse d'abondance différentielle n'est faite : décision de JF du 2026-10-06 (décision 6).]")
# 2. financement
p116 = find1("the contract that also funded Guivier et al. [31]")
replace_once(p116, "est retirée avec elle — à valider par JF.", "est retirée avec elle — retrait validé par JF le 2026-10-06.")
# 4. limites : taille
lim = find1("Four limits bound these conclusions.")
replace_once(lim, "Four limits bound these conclusions.", "Five limits bound these conclusions.")
anchor_txt = "is not resolved by these data."
hits = [r for r in lim.runs if anchor_txt in r.text]
assert len(hits) == 1 and hits[0].text.count(anchor_txt) == 1
r0 = hits[0]; cut = r0.text.index(anchor_txt) + len(anchor_txt); tail = r0.text[cut:]
r0.text = r0.text[:cut]
segs = [(" Body size also differed between genotypic categories (mean total length 24.4 cm in ", False), ("C. nasus", True),
        (", 19.4 cm in hybrids and 16.0 cm in ", False), ("P. toxostoma", True),
        ("; Kruskal–Wallis p = 0.004), and age was not estimated, so that category contrasts may include effects of size or "
         "age that this design cannot separate from those of ancestry.", False)]
if tail: segs.append((tail, None))
prev = r0._r
for t, it in segs:
    el = copy.deepcopy(r0._r); prev.addnext(el); prev = el
    rr = Run(el, lim); rr.text = t; rr.italic = True if it else None
txt = "\n".join(p.text for p in A.paragraphs)
assert txt.count("mean total length 24.4 cm in C. nasus, 19.4 cm in hybrids and 16.0 cm in P. toxostoma; Kruskal–Wallis p = 0.004") == 1
assert "à trancher par JF" not in p52.text and "à valider par JF" not in p116.text
A.save(f"{OUTD}/Article.docx")
print("Article.docx", hashlib.md5(open(f"{OUTD}/Article.docx", "rb").read()).hexdigest())
