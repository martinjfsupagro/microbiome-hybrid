#!/usr/bin/env python3
"""72-article_soumission.py — copie de soumission de l'article pour Animal Microbiome (article type Research), dérivée
de la version de travail docs/manuscrit/Article.docx sans la modifier (décision de JF du 2026-10-06 : « copie de
soumission sans notes », un Additional file par élément du supplément).

Transformations (consignes « Prepare your manuscript » et « Research », lues le 2026-10-06) :
  1. notes éditoriales grises retirées (runs de couleur 888888) ; paragraphes devenus vides supprimés, sauf les
     sections obligatoires des Déclarations, complétées par des marqueurs [TO COMPLETE: ...] ;
  2. couleurs de texte remises à « automatique » ;
  3. page de titre : auteurs, adresses, courriels, auteur correspondant en marqueurs [TO COMPLETE] ;
  4. appels « Supplementary Table/Figure/Note Sn » -> « Additional file n » (soumission_commun.conv, ordre de
     première citation) ; section « Additional files » en fin de document (nom, format, titre, description) ;
  5. URL dans le texte des Méthodes (§9) remplacée par un renvoi à « Availability of data and materials », où le
     dépôt est déjà la référence [51] ; lien ENA ajouté ;
  6. « 16S » -> « 16S rRNA gene » (terminologie demandée par la revue), hors titres de références ;
  7. noms d'espèces et de taxons microbiens en italique ;
  8. double interligne, numéros de ligne continus, numéro de page en pied de page, aucun saut de page.
Contrôles en fin de script (assertions) : voir la fin du fichier.
Usage : 72-article_soumission.py <Article.docx de travail> <Supplementary_Data.docx de travail> <sortie.docx>
"""
import sys, re, copy, hashlib
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
sys.path.insert(0, __file__.rsplit("/", 1)[0])
import soumission_commun as SC

ART, SUP, OUT = sys.argv[1:4]
assert hashlib.md5(open(ART, "rb").read()).hexdigest() == "45f528985e2917a93649807822e7e874", "Article : md5 inattendu"
assert hashlib.md5(open(SUP, "rb").read()).hexdigest() == "e050a3545fdba5e6ec9b014ed707b1a8", "Supplément : md5 inattendu (attendu : sortie du script 74)"
A = docx.Document(ART)


def para_index(txt):
    k = [i for i, p in enumerate(A.paragraphs) if p.text.strip() == txt]
    assert len(k) == 1, (txt, k); return k[0]


def replace_once(par, old, new):
    hits = [r for r in par.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (old[:60], len(hits))
    hits[0].text = hits[0].text.replace(old, new)

# ---------------------------------------------------------------- 1. notes grises
had_note = []                                               # références aux éléments (id() instable sous lxml)
for i, p in enumerate(A.paragraphs):
    for r in list(p.runs):
        if SC.est_note(r):
            if p._p not in had_note: had_note.append(p._p)
            r._r.getparent().remove(r._r)
    if p._p in had_note and p.runs:                      # espace laissé avant la note
        last = [r for r in p.runs if r.text][-1:] 
        if last: last[0].text = last[0].text.rstrip()
KEEP_EMPTY = {"Authors' contributions"}
heads = {p.text.strip(): i for i, p in enumerate(A.paragraphs) if p.style.name.startswith("Heading")}
keep_ids = []
for h in KEEP_EMPTY:
    keep_ids.append(A.paragraphs[heads[h] + 1]._p)
removed = 0
for p in list(A.paragraphs):
    if p._p in had_note and not p.text.strip() and p._p not in keep_ids:
        p._p.getparent().remove(p._p); removed += 1
assert not any(SC.est_note(r) for p in A.paragraphs for r in p.runs)
# ---------------------------------------------------------------- 2. couleurs
for p in A.paragraphs:
    for r in p.runs:
        if SC.couleur(r): r.font.color.rgb = None
# ---------------------------------------------------------------- 3. page de titre
P = A.paragraphs
assert P[0].text == "Host-associated bacterial microbiota of hybridising cyprinid fishes"
anchor = P[0]._p
for t in reversed(["[TO COMPLETE: full names of all authors, with affiliation numbers]",
                   "[TO COMPLETE: institutional addresses of all authors]",
                   "[TO COMPLETE: e-mail addresses of all authors]",
                   "Corresponding author: [TO COMPLETE: name and e-mail address]"]):
    el = OxmlElement("w:p"); anchor.addnext(el)
    np_ = Paragraph(el, P[0]._parent); np_.style = A.styles["Normal"]; np_.add_run(t)
# ---------------------------------------------------------------- 4. appels
iref = para_index("References"); ifl = para_index("Figure legends")
for i, p in enumerate(A.paragraphs):
    if iref < i < ifl: continue                              # liste de références intacte
    SC.conv_paragraphe(p)
# ---------------------------------------------------------------- 5. URL et disponibilité
i64 = [i for i, p in enumerate(A.paragraphs) if "The code is publicly available at https://github.com/martinjfsupagro/microbiome-hybrid." in p.text]
assert len(i64) == 1
replace_once(A.paragraphs[i64[0]], "The code is publicly available at https://github.com/martinjfsupagro/microbiome-hybrid.",
             "The code is publicly available (see Availability of data and materials).")
heads = {p.text.strip(): i for i, p in enumerate(A.paragraphs) if p.style.name.startswith("Heading")}
pav = A.paragraphs[heads["Availability of data and materials"] + 1]
replace_once(pav, "in the European Nucleotide Archive under accession PRJEB124417.",
             "in the European Nucleotide Archive under accession PRJEB124417, https://www.ebi.ac.uk/ena/browser/view/PRJEB124417.")
pav.add_run(" [TO COMPLETE: DOI of the archived version of the code (Zenodo), to be cited in the reference list.]")
pfu = A.paragraphs[heads["Funding"] + 1]
pfu.add_run(" [TO COMPLETE: role of the funding bodies in the design of the study, in the collection, analysis and "
            "interpretation of data and in writing the manuscript.]")
pco = A.paragraphs[heads["Authors' contributions"] + 1]
assert not pco.text.strip(), repr(pco.text[:120])
pco.add_run("[TO COMPLETE: contribution of each author, by initials; e.g. “XX designed the study … All authors read and "
            "approved the final manuscript.”]")
pac = A.paragraphs[heads["Acknowledgements"] + 1]
pac.add_run(" [TO COMPLETE: role of Benjamin Hérodet; other acknowledgements.]")
# ---------------------------------------------------------------- 6. 16S rRNA gene
for old, new in SC.TERMES[:4]:
    hits = [p for i, p in enumerate(A.paragraphs) if old in p.text and not (iref < i < ifl)]
    assert len(hits) == 1, (old, len(hits)); replace_once(hits[0], old, new)
# ---------------------------------------------------------------- 7. italiques
iref = para_index("References"); ifl = para_index("Figure legends")
for i, p in enumerate(A.paragraphs):
    if iref < i < ifl: continue
    for r in list(p.runs): SC.italiciser_run(p, r)
# ---------------------------------------------------------------- 4bis. section Additional files
S = docx.Document(SUP); body = list(S.element.body.iterchildren())
def ptext(el):
    p = Paragraph(el, S)
    return "".join(r.text for r in p.runs if not SC.est_note(r)).strip()
body, ELEMS = SC.elements_supplement(S)
starts = [(k0, lab, tit) for k0, k1, lab, tit in ELEMS]
import csv
COR = {r["element"]: r for r in csv.DictReader(open(sys.argv[4] if len(sys.argv) > 4 else "docs/soumission/correspondance_additional_files.tsv"), delimiter="\t")}
NOTE_DESC = {"N1": "Confirmation on the deposited data of the nested replication structure (sequencing run within library "
                   "preparation) and its consequence for the ENA record.",
             "N2": "Full parameters and controls of library preparation, sequence processing, quality control, removal of host "
                   "12S rRNA gene reads, decontamination, phylogeny, repeated rarefaction, ENA submission and metadata "
                   "corrections, summarised in the Methods.",
             "N3": "Reuse of field fish identifiers between sampling campaigns, and the campaign prefix added to the deposited "
                   "host subject identifiers."}
DESC = {}
for j, (k0, lab, _) in enumerate(starts):
    k1 = starts[j + 1][0] if j + 1 < len(starts) else len(body)
    if lab[0] == "N": DESC[lab] = NOTE_DESC[lab]; continue
    if lab[0] == "F":
        cap = ptext(body[k0 + 1]); DESC[lab] = re.sub(r"^Figure S\d+\.\s*", "", cap); continue
    txt = [ptext(body[k]) for k in range(k0 + 1, k1) if body[k].tag.endswith("}p")]
    DESC[lab] = " ".join(t for t in txt if t)
DESC["T7"] = SC.texte_T7(DESC["T7"])
assert "ENA_accessions_durance.tsv" not in DESC["T7"] and "excerpt" not in DESC["T7"], DESC["T7"][:300]
for lab in DESC: DESC[lab] = SC.terminologie(SC.conv(DESC[lab]))
assert [SC.HIST_COMPTE[a] for a, _ in SC.HISTORIQUE] == [1, 1, 0, 1, 1], SC.HIST_COMPTE  # AF 1 (x2), 12, 17 ; la description de l'AF 6 est rédigée à part
last = A.paragraphs[-1]
h = A.add_paragraph("Additional files", style=A.paragraphs[para_index("Figure legends")].style)
for lab in SC.ORDER:
    n = SC.AF[lab]
    p = A.add_paragraph(); r = p.add_run(f"Additional file {n}"); r.bold = True
    p.add_run(f" (.{SC.FMT[lab[0]]}). ")
    for t, it in SC.segments_italiques(COR[lab]["titre_court"] + "."):
        rr = p.add_run(t); rr.bold = True; rr.italic = True if it else None
    p2 = A.add_paragraph()
    for t, it in SC.segments_italiques(DESC[lab]):
        rr = p2.add_run(t); rr.italic = True if it else None
# ---------------------------------------------------------------- 8. mise en page
for st in A.styles:
    try:
        if st.type == 1: st.paragraph_format.line_spacing = 2.0
    except Exception: pass
for p in A.paragraphs:
    p.paragraph_format.line_spacing = 2.0; p.paragraph_format.page_break_before = False
    for b in p._p.iter(qn("w:br")):
        assert b.get(qn("w:type")) != "page"
for sec in A.sections:
    sp = sec._sectPr
    for old in sp.findall(qn("w:lnNumType")): sp.remove(old)
    ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
    pg = sp.find(qn("w:pgSz")); (pg.addnext(ln) if pg is not None else sp.append(ln))
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    for r in list(fp.runs): r._r.getparent().remove(r._r)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); run._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)
A.save(OUT)

# ---------------------------------------------------------------- contrôles
B = docx.Document(OUT); T = [p.text for p in B.paragraphs]
iref = T.index("References"); ifl = T.index("Figure legends"); iaf = T.index("Additional files")
corps = "\n".join(T[:iref] + T[ifl:iaf])
assert not SC.residuel(corps), SC.residuel(corps)
brack = [m.group(0) for m in re.finditer(r"\[[^\]]*\]", corps) if not re.fullmatch(r"\[\d+(?:\s*[–,-]\s*\d+)*\]", m.group(0))]
assert all(b.startswith("[TO COMPLETE") for b in brack), [b for b in brack if not b.startswith("[TO COMPLETE")][:5]
cit = []
for m in re.finditer(r"Additional files? ((?:\d+(?:, | and )?)+)", corps):
    for n in re.findall(r"\d+", m.group(1)):
        if int(n) not in cit: cit.append(int(n))
assert cit == list(range(1, 19)), cit
assert "Additional file 12 and Additional file 13" not in corps
assert all(B.paragraphs[i].style.name == "Normal" for i in range(1, 5))
assert not any(SC.couleur(r) for p in B.paragraphs for r in p.runs)
roman = [r.text for i, p in enumerate(B.paragraphs) if not (iref < i < ifl) for r in p.runs if not r.italic and SC.ITAL.search(r.text)]
assert not roman, roman[:5]
ln = B.sections[0]._sectPr.find(qn("w:lnNumType")); assert ln is not None
assert all(p.paragraph_format.line_spacing == 2.0 for p in B.paragraphs)
hors_ref = "\n".join(T[:iref] + T[ifl:])
assert not SC.seize_s_isole(hors_ref), SC.seize_s_isole(hors_ref)
print(f"{OUT} | md5 {hashlib.md5(open(OUT, 'rb').read()).hexdigest()} | paragraphes retirés {removed} | marqueurs {len(brack)} | "
      f"Additional files cités dans l'ordre 1-18")
