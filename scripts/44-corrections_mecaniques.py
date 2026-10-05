"""44-corrections_mecaniques.py — corrections du 2026-10-05, appliquées à l'identique aux deux versions de l'article
(Article.docx, Article_resserre.docx) et du supplément (Supplementary_Data.docx, Supplementary_Data_resserre.docx).
1. Supplément : titres des Tables S6, S8, S9, S10 en style « Heading 1 ».
2. Figures S renumérotées dans l'ordre de citation (ancienne S2 -> S1, ancienne S1 -> S2) ; la nouvelle Figure S1
   (compromis profondeur / fidélité / sélection) remplace l'ancienne image, en français et dont le panneau b portait des
   valeurs d'une classification périmée ; elle est placée après la Table S6. Médias orphelins retirés.
3. Références : noms de revues abrégés selon le catalogue NLM (MedlineTA, recherche par ISSN).
4. Figure 4 : légende alignée sur les nouveaux titres de panneaux.
5. §8.6 : les écarts de rétention cités sont ceux du run durance1 (script 43) ; précision ajoutée + note grise.
   Discussion : Camper et al. comparent les systèmes publiés dans deux versions de l'indice (Tables 1 et 3), pas une.
   Note ENA « A TRANCHER » périmée remplacée.
Entrées gardées par md5 ; sorties écrites à côté (suffixe inchangé, mêmes noms) dans le répertoire OUT."""
import docx, re, hashlib, json, os, sys, copy
from docx.shared import RGBColor, Pt, Emu
from docx.oxml.ns import qn

M = os.environ.get("MS_DIR", "docs/manuscrit/")
OUT = os.environ.get("OUT_DIR", M)
FIG_S1 = os.environ.get("FIG_S1", M + "figures/fig_S1_depth_tradeoff.png")
NLM = json.load(open(os.environ.get("NLM_JSON", "docs/biblio/nlm_abreviations.json")))
IN = {"Article.docx": "68caaeb7c55ec2716a7fa36010d898dd", "Article_resserre.docx": "965b24f60fc9f4b1445f69e74cfeb759",
      "Supplementary_Data.docx": "3b6f95e4f6ac72b5dc773f53c061d483",
      "Supplementary_Data_resserre.docx": "4e38874db427458aa184608347c92723"}
for f, h in IN.items():
    assert hashlib.md5(open(M + f, "rb").read()).hexdigest() == h, f"md5 inattendu : {f}"
BLUE = RGBColor(0x1F, 0x4E, 0x79); GREY = RGBColor(0x88, 0x88, 0x88)
LOG = []

def grey_run(p, text, before=None):
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY
    if before is not None: before.addprevious(r._r)
    return r

def sub_run(p, old, new, blue=True, count=1):
    hit = [r for r in p.runs if old in r.text]
    assert len(hit) == 1, (old, len(hit))
    r = hit[0]
    if not blue:
        r.text = r.text.replace(old, new, count); return r
    a, b = r.text.split(old, 1)
    r.text = a
    nr = copy.deepcopy(r._r); r._r.addnext(nr)
    nr2 = copy.deepcopy(r._r); nr.addnext(nr2)
    from docx.text.run import Run
    R1, R2 = Run(nr, p), Run(nr2, p)
    R1.text = new; R1.font.color.rgb = BLUE
    R2.text = b
    return R1

def swap_fig_s(doc):
    n = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if re.search(r"Figures? S[12]\b", r.text):
                t = re.sub(r"(Figures?) S1\b", r"\1 S§§", r.text)
                t = re.sub(r"(Figures?) S2\b", r"\1 S1", t)
                r.text = t.replace("S§§", "S2"); n += 1
    return n

def one(doc, s):
    hits = [p for p in doc.paragraphs if s in p.text]
    assert len(hits) == 1, (s, len(hits)); return hits[0]

# --------------------------------------------------------------------------- articles
for f in ("Article.docx", "Article_resserre.docx"):
    A = docx.Document(M + f)
    n = swap_fig_s(A); LOG.append(f"{f}: renvois Figure S1/S2 permutés dans {n} segments")
    # §8.6 : run durance1
    p = one(A, "in the caudal fin 50 % of intermediate hybrids")
    sub_run(p, "in the caudal fin 50 % of intermediate hybrids", "in the caudal fin, in sequencing run durance1, 50 % of intermediate hybrids")
    grey_run(p, " [Écarts de rétention : valeurs du run durance1 (scripts/43-retention_profondeur.py, "
                "results/depth_agreement/retention_par_categorie.tsv, 2026-10-05) ; aucun script ne les produisait auparavant. "
                "Sur les trois runs : caudale à 3 000 lectures 29,0–36,4 points (passe 1) et 14,7–19,8 (passe 2) ; à 500 lectures "
                "2,6–3,8 et 3,4–5,7 ; midgut à 500 lectures 5,7–14,0 et 7,8–8,9 (runs confondus : 32,3 / 17,2 ; 3,4 / 4,5 ; 8,5 / 8,2). "
                "Le midgut de la passe 1 dépend du run (14,0 dans durance1, 5,7 dans les deux autres). Rapporter la valeur d'un run "
                "ou l'étendue : à trancher par JF.]")
    # Figure 4 : légende
    p = one(A, "Values below 0: hybrids less dispersed than the parental mean.")
    sub_run(p, "Values below 0: hybrids less dispersed than the parental mean.",
            "Values below 0: hybrids less dispersed than the mean of the two parents; panel titles give the number of tests "
            "in which hybrids were less dispersed than both parents, a stricter condition.")
    # Discussion : Camper et al., deux versions comparées
    p = one(A, "the form in which the framework's authors compared published systems [14]")
    sub_run(p, "the form in which the framework's authors compared published systems [14]",
            "one of the two forms in which the framework's authors compared published systems [14]")
    # Références : abréviations NLM
    k = 0; sec = None
    for p in A.paragraphs:
        if p.style.name == "Title": sec = p.text; continue
        if sec != "References": continue
        m = re.match(r"(\d+)\. ", p.text)
        if not m or m.group(1) not in NLM: continue
        old, new = NLM[m.group(1)]
        if old == new: continue
        assert len(p.runs) == 1 and f". {old}. " in p.runs[0].text, (f, m.group(1))
        p.runs[0].text = p.runs[0].text.replace(f". {old}. ", f". {new}. ", 1); k += 1
    LOG.append(f"{f}: {k} noms de revues abrégés")
    p = one(A, "noms de revues non abrégés (à abréger selon la norme NLM avant soumission)")
    for r in p.runs:
        if "noms de revues non abrégés" in r.text:
            r.text = r.text.replace("noms de revues non abrégés (à abréger selon la norme NLM avant soumission)",
                                    "noms de revues abrégés le 2026-10-05 selon le catalogue NLM (MedlineTA, recherche par ISSN ; "
                                    "scripts/44-corrections_mecaniques.py) ; préprint, thèse, logiciels et dépôt non abrégés")
    A.save(OUT + f)

# --------------------------------------------------------------------------- suppléments
NEW_LEG = ("Figure S1. Trade-off between rarefaction depth, metric fidelity and selection bias. (a) Pearson correlation between "
           "the dissimilarity matrices computed at each depth and at 3,000 reads, on the fixed set of 1,784 samples retained at "
           "3,000 reads, so that only depth varies. (b) Loss of observed richness relative to 3,000 reads, same samples. "
           "(c, d) Retention gap: percentage of P. toxostoma samples minus percentage of hybrid samples reaching each depth, "
           "in the caudal fin (c) and the midgut (d), for pass 1 (20 intermediate hybrids) and pass 2 (all 42 hybrids). "
           "Lines: sequencing run durance1 (values quoted in Materials and Methods §8.6); shaded bands: range across the three runs.")
for f in ("Supplementary_Data.docx", "Supplementary_Data_resserre.docx"):
    S = docx.Document(M + f); P = list(S.paragraphs)
    # 1. styles des titres de tables
    for p in P:
        if re.match(r"Table S(6|8|9|10)\. ", p.text): p.style = S.styles["Heading 1"]
    # 2. bloc de l'ancienne Figure S2 (compromis) : image + légende, déplacés après la note de la Table S6
    leg = one(S, "Figure S2. Trade-off between rarefaction depth")
    i = [q._p for q in P].index(leg._p); img = P[i - 1]; blank_before = P[i - 2]
    assert img._p.findall(".//" + qn("a:blip")) and not blank_before.text.strip()
    t6 = one(S, "Table S6. Agreement of diversity metrics"); j = [q._p for q in P].index(t6._p)
    note6 = P[j + 1]; blank6 = P[j + 2]; assert not blank6.text.strip()
    blank6._p.addnext(img._p); img._p.addnext(leg._p); leg._p.addnext(blank_before._p)
    # image : nouveau fichier, rapport largeur/hauteur conservé à largeur constante
    rid = img._p.findall(".//" + qn("a:blip"))[0].get(qn("r:embed"))
    part = S.part.related_parts[rid]
    from PIL import Image
    data = open(FIG_S1, "rb").read(); w, h = Image.open(FIG_S1).size
    part._blob = data
    ext = img._p.findall(".//" + qn("wp:extent"))[0]; cx = int(ext.get("cx")); cy = int(cx * h / w)
    ext.set("cy", str(cy))
    for e in img._p.findall(".//" + qn("a:ext")): e.set("cy", str(cy))
    # nouvelle légende (remplace l'ancienne)
    for r in list(leg.runs): r._r.getparent().remove(r._r)
    r = leg.add_run("Figure S1."); r.bold = True
    body = NEW_LEG[len("Figure S1."):]
    for k2, seg in enumerate(re.split(r"(P\. toxostoma)", body)):
        if seg:
            rr = leg.add_run(seg); rr.font.color.rgb = BLUE
            if k2 % 2: rr.italic = True
    grey_run(leg, " [Figure régénérée le 2026-10-05 (ex-Figure S2) : l'ancienne version était en français et son panneau b "
                  "portait des valeurs recopiées d'une trace antérieure à la reclassification de septembre (30,4 points à 3 000 "
                  "lectures contre 36,4 recalculés) ; panneaux c et d calculés par scripts/43-retention_profondeur.py.]")
    # renvois Figure S1 <-> S2 dans tout le supplément (dont les deux légendes)
    # la nouvelle légende commence déjà par « Figure S1. » : protéger le premier run avant la permutation
    first = leg.runs[0]; first.text = "Figure S§1."
    n = swap_fig_s(S); first.text = "Figure S1."
    LOG.append(f"{f}: renvois Figure S1/S2 permutés dans {n} segments")
    # §8.6 recopié verbatim en Note S2 (version resserrée) : même précision
    hits = [p for p in S.paragraphs if "in the caudal fin 50 % of intermediate hybrids" in p.text]
    for p in hits:
        sub_run(p, "in the caudal fin 50 % of intermediate hybrids", "in the caudal fin, in sequencing run durance1, 50 % of intermediate hybrids")
    LOG.append(f"{f}: précision durance1 ajoutée dans {len(hits)} paragraphe(s)")
    # 3. note ENA périmée
    p = one(S, "A TRANCHER : rendre 'host subject id'")
    for rr in list(p.runs):
        if "A TRANCHER" in rr.text: rr._r.getparent().remove(rr._r)
    grey_run(p, "[Décidé par JF le 2026-10-03 : le « host subject id » sera préfixé par la campagne (14Bue1004 / 15Bue1004) "
                "dans la même correction MODIFY que les dates de Pertuis ; soumission à lancer par JF (identifiants Webin). "
                "Réécrire cette note une fois la correction vérifiée sur l'archive.]")
    # 4. médias orphelins
    used = {b.get(qn("r:embed")) for b in S.element.body.iter(qn("a:blip"))}
    drop = [k3 for k3, rel in S.part.rels.items() if rel.reltype.endswith("/image") and k3 not in used]
    for k3 in drop:
        try: del S.part.rels[k3]
        except TypeError: S.part.rels._rels.pop(k3)
    LOG.append(f"{f}: {len(drop)} médias orphelins retirés")
    S.save(OUT + f)
print("\n".join(LOG))
