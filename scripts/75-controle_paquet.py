#!/usr/bin/env python3
"""75-controle_paquet.py — contrôle du paquet de soumission Animal Microbiome (consignes lues le 2026-10-06 :
guidelines/prep_manuscript.txt, research.txt, prep_supporting.txt).

Usage : 75-controle_paquet.py <Manuscript.docx> <dossier figures Fig1-5> <dossier additional_files> <correspondance.tsv> <rapport.tsv>
Chaque contrôle écrit une ligne : id, objet, règle, valeur, statut (OK / ECART / A COMPLETER). Code de sortie 1 si ECART.
"""
import sys, os, re, csv, zipfile
import docx
from docx.oxml.ns import qn
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import soumission_commun as SC

MS, FIGD, AFD, CORF, OUT = sys.argv[1:6]
R = []
def chk(i, objet, regle, valeur, ok, statut_si_non="ECART"):
    R.append((i, objet, regle, str(valeur), "OK" if ok else statut_si_non))

D = docx.Document(MS); P = D.paragraphs; T = [p.text for p in P]
titres = {p.text.strip(): i for i, p in enumerate(P) if p.style.name == "Title"}
iref, ifl, iaf = titres["References"], titres["Figure legends"], titres["Additional files"]
corps = "\n".join(T[:iref])
# --- M1 couleurs
cols = sum(1 for p in P for r in p.runs if SC.couleur(r))
chg = sum(1 for p in P for r in p.runs if SC.est_note(r))
chk("M1", "manuscrit", "aucune note grise ni couleur", f"notes {chg}, runs colorés {cols}", cols == 0 and chg == 0)
# --- M2 crochets
cro = [m.group(0) for m in re.finditer(r"\[[^\]]*\]", "\n".join(T)) if not re.fullmatch(r"\[\d+(?:\s*[,–-]\s*\d+)*\]", m.group(0))]
tc = [x for x in cro if x.startswith("[TO COMPLETE")]; autres = [x for x in cro if not x.startswith("[TO COMPLETE")]
chk("M2a", "manuscrit", "crochets hors références : seulement [TO COMPLETE]", "; ".join(autres[:6]) or "aucun", not autres)
chk("M2b", "manuscrit", "marqueurs [TO COMPLETE] à remplir avant soumission", len(tc), len(tc) == 0, "A COMPLETER")
# --- M3 Additional files
prem = []
for m in re.finditer(r"Additional files? ((?:\d+)(?:(?:, | and )\d+)*)", corps):
    for n in map(int, re.findall(r"\d+", m.group(1))):
        if n not in prem: prem.append(n)
chk("M3a", "manuscrit", "Additional files cités dans le corps, 1re citation dans l'ordre 1-18", prem, prem == list(range(1, 19)))
lst = [int(m.group(1)) for t in T[iaf + 1:] for m in [re.match(r"Additional file (\d+) \(\.(xlsx|docx|pdf)\)\. ", t)] if m]
chk("M3b", "manuscrit", "section Additional files : 1-18 avec format", lst, lst == list(range(1, 19)))
res = SC.residuel("\n".join(T)) + re.findall(r"Supplementary (?:Table|Figure|Note|Data)", "\n".join(T))
chk("M3c", "manuscrit", "aucun renvoi Supplementary / S résiduel", res[:5] or "aucun", not res)
# --- M4 figures
fp = []
for m in re.finditer(r"Fig(?:ure)?s?\.?\s+(\d)(?:[a-z,–-]*\s*(?:and|–|-)\s*(\d))?", corps):
    for g in m.groups():
        if g and int(g) not in fp: fp.append(int(g))
chk("M4a", "manuscrit", "Figures 1-5 citées dans l'ordre", fp, fp == [1, 2, 3, 4, 5])
leg = [t for t in T[ifl + 1:iaf] if re.match(r"Fig(?:ure)?\.? ?\d", t)]
for t in leg:
    n = re.match(r"Fig(?:ure)?\.? ?(\d)", t).group(1)
    corpsleg = re.sub(r"^Fig(?:ure)?\.? ?\d+\.?\s*", "", t)
    titre, reste = (corpsleg.split(". ", 1) + [""])[:2]
    chk(f"M4b{n}", f"Figure {n}", "titre ≤ 15 mots", len(titre.split()), len(titre.split()) <= 15)
    chk(f"M4c{n}", f"Figure {n}", "légende ≤ 300 mots", len(reste.split()), len(reste.split()) <= 300)
chk("M4d", "manuscrit", "5 légendes de figures", len(leg), len(leg) == 5)
# --- M5 résumé et mots-clés
ia = titres["Abstract"]; ib = titres["Background"]
ab = [t for t in T[ia + 1:ib] if t.strip()]
kw = [t for t in ab if t.lower().startswith("keywords")]
abs_ = [t for t in ab if not t.lower().startswith("keywords")]
nmots = sum(len(t.split()) for t in abs_)
chk("M5a", "résumé", "≤ 350 mots (intertitres compris)", nmots, nmots <= 350)
secs = [re.match(r"(Background|Methods|Results|Conclusions?)[:.]?", t).group(1) for t in abs_ if re.match(r"(Background|Methods|Results|Conclusions?)[:.]?", t)]
chk("M5b", "résumé", "sections Background, Results, Conclusions", secs, all(s in " ".join(secs) for s in ("Background", "Results", "Conclusion")))
chk("M5c", "résumé", "aucune référence citée", len(re.findall(r"\[\d", " ".join(abs_))), not re.search(r"\[\d", " ".join(abs_)))
nk = len(re.split(r"[;,]\s*", kw[0].split(":", 1)[1].strip())) if kw else 0
chk("M5d", "mots-clés", "3 à 10", nk, 3 <= nk <= 10)
# --- M6 Declarations
idc = titres["Declarations"]; sous = [t.strip() for t in T[idc + 1:iref] if P[T.index(t, idc)].style.name.startswith("Heading")]
for s in ["Ethics approval and consent to participate", "Consent for publication", "Availability of data and materials",
          "Competing interests", "Funding", "Authors' contributions", "Acknowledgements"]:
    ok = any(x.replace("’", "'").startswith(s.replace("material", "material")) or x.replace("’", "'") == s for x in sous)
    chk("M6", "Declarations", s, "présent" if ok else "absent", ok)
for s in ["Background", "Methods", "Results", "Discussion", "Conclusions", "Abbreviations"]:
    chk("M6s", "sections", s, "présent" if s in titres else "absent", s in titres)
# --- M7 mise en page
from docx.shared import Pt
def interligne(p):
    pf = p.paragraph_format; st = p.style
    while pf.line_spacing is None and st is not None: pf = st.paragraph_format; st = st.base_style
    return pf.line_spacing
corpsP = [p for p in P[:iref] if p.text.strip() and p.style.name != "Title" and not p.style.name.startswith("Heading")]
dbl = sum(1 for p in corpsP if interligne(p) in (2, 2.0)); chk("M7a", "manuscrit", "double interligne (corps)", f"{dbl}/{len(corpsP)}", dbl == len(corpsP))
sp = D.sections[0]._sectPr
chk("M7b", "manuscrit", "numéros de ligne continus", sp.find(qn("w:lnNumType")) is not None and sp.find(qn("w:lnNumType")).get(qn("w:restart")) == "continuous",
    sp.find(qn("w:lnNumType")) is not None and sp.find(qn("w:lnNumType")).get(qn("w:restart")) == "continuous")
foot = D.sections[0].footer._element.xml
chk("M7c", "manuscrit", "numéro de page (champ PAGE)", "PAGE" in foot, "PAGE" in foot)
# --- M8 italiques, terminologie
rom = [(i, r.text) for i, p in enumerate(P) if not (iref < i < ifl) for r in p.runs if SC.ITAL.search(r.text) and not r.italic]
chk("M8a", "manuscrit", "espèces et taxons en italique (hors références)", len(rom), not rom)
chk("M8b", "manuscrit", "« 16S » suivi de « rRNA » (hors références)", SC.seize_s_isole("\n".join(T[:iref] + T[ifl:]))[:3] or "aucun", not SC.seize_s_isole("\n".join(T[:iref] + T[ifl:])))
chk("M8c", "manuscrit", "intitulé « Methods » (pas « Materials and Methods »)", corps.count("Materials and Methods"), "Materials and Methods" not in "\n".join(T))
chk("M8d", "manuscrit", "aucune table dans le manuscrit (toutes en Additional files)", len(D.tables), len(D.tables) == 0)
# --- F figures principales
for n in range(1, 6):
    f = os.path.join(FIGD, f"Fig{n}.png"); ok = os.path.exists(f)
    if ok:
        im = Image.open(f); taille = os.path.getsize(f); dpi = im.info.get("dpi", (0, 0))[0]
        larg_mm = im.size[0] / dpi * 25.4 if dpi else None
        chk(f"F{n}a", f"Fig{n}.png", "≤ 10 Mo", f"{taille/1e6:.2f} Mo", taille <= 10e6)
        chk(f"F{n}b", f"Fig{n}.png", "≥ 300 dpi à 170 mm", f"{im.size[0]} px ({im.size[0]/(170/25.4):.0f} dpi à 170 mm)", im.size[0] / (170 / 25.4) >= 295)
    else: chk(f"F{n}a", f"Fig{n}.png", "présent", "absent", False)
# --- A additional files
COR = list(csv.DictReader(open(CORF), delimiter="\t"))
noms = sorted(os.listdir(AFD))
for r in COR:
    n = int(r["additional_file"]); f = os.path.join(AFD, r["fichier"]); ok = os.path.exists(f)
    chk(f"A{n}a", r["fichier"], "présent, nom « Additional_file_n.ext »", ok, ok and re.fullmatch(rf"Additional_file_{n}\.(xlsx|docx|pdf)", r["fichier"]) is not None)
    if not ok: continue
    chk(f"A{n}b", r["fichier"], "≤ 10 Mo", f"{os.path.getsize(f)/1e6:.3f} Mo", os.path.getsize(f) <= 10e6)
    chk(f"A{n}c", r["fichier"], "titre ≤ 15 mots", len(r["titre_court"].split()), len(r["titre_court"].split()) <= 15)
    if f.endswith(".pdf"):
        # objets descripteurs (« /Type /FontDescriptor ») et non les renvois « /FontDescriptor n 0 R » des polices CID
        b = open(f, "rb").read(); objs = re.findall(rb"\d+ 0 obj\s*(<<.*?>>)\s*endobj", b, re.S)
        fds = [d for d in objs if re.search(rb"/Type\s*/FontDescriptor", d)]; sans = [d for d in fds if not re.search(rb"/FontFile[23]?", d)]
        t3 = sum(1 for d in objs if re.search(rb"/Subtype\s*/Type3", d))
        chk(f"A{n}d", r["fichier"], "polices incorporées", f"descripteurs {len(fds)}, sans fichier de police {len(sans)}, Type3 {t3}", not sans)
extra = [x for x in noms if x.startswith("Additional_file_") and x not in {r["fichier"] for r in COR}]
chk("A0", "additional_files", "aucun fichier hors correspondance", extra or "aucun", not extra)
tot = sum(os.path.getsize(os.path.join(AFD, r["fichier"])) for r in COR)
chk("A00", "additional_files", "volume total", f"{tot/1e6:.2f} Mo", True)

with open(OUT, "w") as fh:
    fh.write("id\tobjet\tregle\tvaleur\tstatut\n")
    for r in R: fh.write("\t".join(r) + "\n")
from collections import Counter
print(dict(Counter(r[4] for r in R)))
for r in R:
    if r[4] != "OK": print(*r, sep=" | ")
sys.exit(1 if any(r[4] == "ECART" for r in R) else 0)
