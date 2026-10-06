#!/usr/bin/env python3
"""71-ajouts_largue.py — résultats du script 70 (nageoire caudale du canal du Largue) dans les versions de travail
(décision de JF du 2026-10-06 : légende de la Figure S4 + phrase des limites).
Chiffres (results/largue_caudale/, scripts/70, commit 52cbc91 ; contrôlés par scripts/40) :
  - part du 12S de l'hôte dans les lectures brutes, médiane groupée sur les trois runs : nageoires caudales du Largue
    (60 échantillons, tous plaque 2 colonnes 06-08) 93 % ; nageoires caudales de Saint-Just dans les mêmes colonnes (12)
    74 % ; autres nageoires caudales (474) 25 % ;
  - écart de rétention Pt − Hy en nageoire caudale à 3 000 lectures, passe 1, étendue sur les runs : 29,0–36,4 points
    avec le Largue, 21,5–30,8 sans ;
  - Largue, durance1 : médiane 240,5 lectures après traitement, aucune >= 3 000.
Formulation : l'excès de 12S suit la position (témoin : Saint-Just dans les mêmes colonnes / ailleurs) ; aucun
mécanisme n'est énoncé.
python-docx requis : exécuté hors cluster, fichiers déposés avec garde md5.
Usage : 71-ajouts_largue.py <dossier_entree> <dossier_sortie>
"""
import sys, copy, hashlib
import docx
from docx.text.run import Run

IN, OUTD = sys.argv[1], sys.argv[2]
MD5 = {"Article.docx": "b2791c95fa5ecd660af46ea47c8bbeca", "Supplementary_Data.docx": "be86753f1f8653bc8842d9f71414790b"}
for f, h in MD5.items():
    assert hashlib.md5(open(f"{IN}/{f}", "rb").read()).hexdigest() == h, f"{f} : md5 inattendu"


def insert_after(par, anchor_run, segs):
    prev = anchor_run._r
    for t, it in segs:
        el = copy.deepcopy(anchor_run._r); prev.addnext(el); prev = el
        r = Run(el, par); r.text = t; r.italic = True if it else None

A = docx.Document(f"{IN}/Article.docx")
lim = [p for p in A.paragraphs if p.text.startswith("Five limits bound these conclusions.")]
assert len(lim) == 1; lim = lim[0]
tail = " in the caudal fin, a selection that is declared but cannot be removed for presence-based metrics."
hits = [r for r in lim.runs if r.text == tail or r.text.startswith(tail)]
assert len(hits) == 1, len(hits)
h = hits[0]; rest = h.text[len(tail):]; h.text = tail
segs = [(" Part of this selection follows sample position rather than genotype: the caudal-fin samples placed in three "
         "columns of one plate, from all 20 fish of the Largue canal and 4 of Saint-Just-d'Ardèche, consisted mostly of host "
         "mitochondrial 12S rRNA gene reads (median 93 % and 74 % of raw reads, against 25 % for the other caudal-fin "
         "samples), and excluding the Largue canal narrows the gap between ", False), ("P. toxostoma", True),
        (" and intermediate hybrids from 29–36 to 22–31 points across runs in pass 1, so that most of it remains.", False)]
if rest: segs.append((rest, False))
insert_after(lim, h, segs)
txt = "\n".join(p.text for p in A.paragraphs)
assert txt.count("(median 93 % and 74 % of raw reads, against 25 % for the other caudal-fin samples)") == 1
A.save(f"{OUTD}/Article.docx")

S = docx.Document(f"{IN}/Supplementary_Data.docx")
cap = [p for p in S.paragraphs if p.text.startswith("Figure S4. Community composition")]
assert len(cap) == 1
old = "No caudal-fin sample of the Largue canal reached 3,000 reads in this run."
hits = [r for r in cap[0].runs if old in r.text]; assert len(hits) == 1
hits[0].text = hits[0].text.replace(old, "Caudal-fin samples of the Largue canal are absent from (b): in this run none reached "
    "3,000 reads (median 240.5), host mitochondrial 12S rRNA gene reads making up a median of 93 % of their raw reads "
    "across the three runs.")
S.save(f"{OUTD}/Supplementary_Data.docx")
for f in MD5:
    print(f, hashlib.md5(open(f"{OUTD}/{f}", "rb").read()).hexdigest())
