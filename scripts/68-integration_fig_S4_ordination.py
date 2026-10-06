#!/usr/bin/env python3
"""68-integration_fig_S4_ordination.py — intègre l'ordination (version 1 : b = trois stations, effet station ; choix de
JF et André, 2026-10-06) comme Figure S4 du supplément et l'appelle dans Results « Composition by genotypic
category », après la part de variance de la station. Renumérotation dans l'ordre de première citation :
  S3 phylums (inchangée) -> S4 ordination (nouvelle) -> S5 PERMDISP (ex S4, ex S3).
Données : scripts/65 (commit 25b87d4) ; tracé docs/manuscrit/figures/src/figure_S4_ordination_src.py (v1).
python-docx requis (absent sur meso) : exécuté hors cluster, fichiers déposés avec garde md5.
Usage : 68-integration_fig_S4_ordination.py <dossier_entree> <dossier_sortie> <fig_S4_ordination.png>
"""
import sys, copy, hashlib, re
import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

IN, OUTD, PNG = sys.argv[1], sys.argv[2], sys.argv[3]
MD5 = {"Article.docx": "4e141d99b3fe5eb7c22f9da32bf29479", "Supplementary_Data.docx": "764bebdc8d1ca2d0c6d3aa5ca8dfd656"}
for f, h in MD5.items():
    assert hashlib.md5(open(f"{IN}/{f}", "rb").read()).hexdigest() == h, f"{f} : md5 inattendu"


def replace_once(par, old, new):
    hits = [r for r in par.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (old[:60], len(hits))
    hits[0].text = hits[0].text.replace(old, new)


def fig_order(text):
    o = []
    for m in re.finditer(r"Figures?\s+S(\d+)", text):
        n = int(m.group(1))
        if n not in o: o.append(n)
    return o

# ------------------------------------------------------------------ article
A = docx.Document(f"{IN}/Article.docx"); P = A.paragraphs
assert P[83].text.strip() == "Composition by genotypic category" and P[84].text.startswith("Pseudomonadota was the most abundant")
assert P[85].text.startswith("Fitted in both orderings") and P[88].text.strip() == "Dispersion"
replace_once(P[85], "(Fig. 3c; Supplementary Table S9). No tissue",
             "(Fig. 3c; Supplementary Table S9). In ordination space, the samples of the three stations carrying all three "
             "categories group by station rather than by category in every tissue (Supplementary Figure S4). No tissue")
k = [i for i, p in enumerate(P) if "Supplementary Table S10 and Figure S4." in p.text]
assert len(k) == 1 and k[0] > 88
replace_once(P[k[0]], "Supplementary Table S10 and Figure S4.", "Supplementary Table S10 and Figure S5.")
replace_once(P[k[0]], " et Figure S4 (numérotée S3 avant le 2026-10-06) du supplément",
             " et Figure S5 (numérotée S3, puis S4, avant le 2026-10-06) du supplément")
TXT = "\n".join(p.text for p in A.paragraphs)
assert fig_order(TXT) == [1, 2, 3, 4, 5], fig_order(TXT)
A.save(f"{OUTD}/Article.docx")

# ------------------------------------------------------------------ supplément
S = docx.Document(f"{IN}/Supplementary_Data.docx"); Q = S.paragraphs
assert Q[77].text.startswith("Figure S3. Phylum-level") and Q[78].text == "" and Q[79].text.startswith("Table S10.")
k5 = [i for i, p in enumerate(Q) if p.text.startswith("Figure S4. Dispersion")]
assert len(k5) == 1
replace_once(Q[k5[0]], "Figure S4. Dispersion", "Figure S5. Dispersion")
CAP = ("Figure S4. Community composition groups by tissue and by station rather than by genotypic category (Jaccard). (a) "
       "Principal coordinates analysis (PCoA) of all samples of sequencing run durance1 (n = 583), coloured by tissue; the first "
       "axis separates the two gut sections from the two external tissues. (b) PCoA computed separately for each tissue on the "
       "samples of the three stations where C. nasus, hybrids and P. toxostoma co-occur (Canal (usine du Largue), Confluence "
       "Buëch-Méouge, Saint-Just-d'Ardèche); colour, genotypic category (hybrids pooled); symbol, station. Samples group by "
       "station rather than by category: in every tissue the station centroids are 2.4 to 10 times further apart than the "
       "category centroids. At Confluence Buëch-Méouge, sampled in 2014 and 2015, samples of the caudal fin, gill and midgut also "
       "separate by campaign; restricting the category tests to station × campaign blocks leaves their results unchanged "
       "(Materials and Methods §8.12, control vi). No caudal-fin sample of the Largue canal reached 3,000 reads in this run. "
       "Distances are mean Jaccard dissimilarities over 400 rarefactions to 3,000 reads, the principal metric of the composition "
       "tests (Fig. 3); percentages are eigenvalues relative to the sum of positive eigenvalues. Points are samples, re-extracted "
       "tissues included, as in the tests. The same ordinations computed on runs durance2 and durance3 agree with durance1 "
       "(Procrustes correlation 0.96–0.99, p = 0.001). The figure is descriptive; category effects are tested in the main text "
       "(Fig. 3).")
anchor = Q[78]._p
blank = copy.deepcopy(Q[78]._p); img = copy.deepcopy(Q[78]._p); cap = copy.deepcopy(Q[77]._p)
for r in img.xpath("./w:r"): img.remove(r)
for r in cap.xpath("./w:r")[1:]: cap.remove(r)
anchor.addnext(blank); anchor.addnext(cap); anchor.addnext(img)
Paragraph(img, Q[78]._parent).add_run().add_picture(PNG, width=Inches(6.0))
Paragraph(cap, Q[78]._parent).runs[0].text = CAP
S2 = S.paragraphs
caps = [p.text[:10] for p in S2 if re.match(r"Figure S\d+\. ", p.text)]
assert caps == ["Figure S1.", "Figure S2.", "Figure S3.", "Figure S4.", "Figure S5."], caps
k4 = [i for i, p in enumerate(S2) if p.text.startswith("Figure S4. Community")][0]
assert S2[k4 - 1]._p.xpath(".//w:drawing") and S2[k4 + 2].text.startswith("Table S10.")
S.save(f"{OUTD}/Supplementary_Data.docx")
for f in MD5:
    print(f, hashlib.md5(open(f"{OUTD}/{f}", "rb").read()).hexdigest())
print("ordre des figures S :", fig_order(TXT), "| légendes :", caps)
