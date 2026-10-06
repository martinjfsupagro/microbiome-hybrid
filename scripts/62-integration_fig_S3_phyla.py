#!/usr/bin/env python3
"""62-integration_fig_S3_phyla.py — intègre la figure de composition en phylums (scripts/58-61) au supplément et à
l'article (demande de JF du 2026-10-06) et renumérote les figures supplémentaires dans l'ordre de première citation.

Place retenue : première phrase de Results « Composition by genotypic category » (paragraphe ajouté avant l'analyse
de variance), entre la Figure S2 (Plate position) et l'ancienne Figure S3 (Dispersion). D'où :
  - nouvelle Figure S3 = composition en phylums (fichier fig_S3_phyla.png, ex fig_S4_phyla.png) ;
  - ancienne Figure S3 (PERMDISP) -> Figure S4 (fichier fig_S4_permdisp.png, ex fig_S3_permdisp.png ; le script 41
    écrit toujours fig_S3_permdisp.png, nom historique).
Supplément : figure insérée après le bloc Figure S2 (avant Table S10), même largeur (6 in) et même mise en forme de
légende (paragraphe copié de la légende de la Figure S2).
python-docx requis (absent sur meso) : exécuté hors cluster, fichiers déposés avec garde md5.
Usage : 62-integration_fig_S3_phyla.py <dossier_entree> <dossier_sortie> <fig_S3_phyla.png>
"""
import sys, copy, hashlib, re
import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

IN, OUTD, PNG = sys.argv[1], sys.argv[2], sys.argv[3]
MD5 = {"Article.docx": "77a1e1efbcb71a417e79a304898262b1", "Supplementary_Data.docx": "8ab1040b050b75eaf42fc9f19179ca14"}
for f, h in MD5.items():
    assert hashlib.md5(open(f"{IN}/{f}", "rb").read()).hexdigest() == h, f"{f} : md5 inattendu"


def replace_once(par, old, new):
    hits = [r for r in par.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (old, len(hits))
    hits[0].text = hits[0].text.replace(old, new)


def fig_order(text):
    order = []
    for m in re.finditer(r"Figures?\s+S(\d+)", text):
        n = int(m.group(1))
        if n not in order: order.append(n)
    return order

# ------------------------------------------------------------------ article
A = docx.Document(f"{IN}/Article.docx"); P = A.paragraphs
assert P[83].text.strip() == "Composition by genotypic category" and P[84].text.startswith("Fitted in both orderings")
assert P[86].text.strip() == "Dispersion"
replace_once(P[88], "Supplementary Table S10 and Figure S3.", "Supplementary Table S10 and Figure S4.")
replace_once(P[88], " et Figure S3 du supplément", " et Figure S4 (numérotée S3 avant le 2026-10-06) du supplément")
new = copy.deepcopy(P[84]._p)
for r in new.xpath("./w:r"): new.remove(r)
P[84]._p.addprevious(new)
np_ = Paragraph(new, P[84]._parent)
for txt, it in [("", False), ("Pseudomonadota", True), (" was the most abundant phylum in 33 of the 36 tissue × station profiles; at the three stations where the three categories co-occur, ", False),
                ("Fusobacteriota", True), (" and ", False), ("Bacillota", True),
                (" were more abundant in both gut sections than in both external tissues in each genotypic category (Supplementary Figure S3).", False)]:
    if txt:
        np_.add_run(txt).italic = True if it else None
TXT = "\n".join(p.text for p in A.paragraphs)
assert fig_order(TXT) == [1, 2, 3, 4], fig_order(TXT)
A.save(f"{OUTD}/Article.docx")

# ------------------------------------------------------------------ supplément
S = docx.Document(f"{IN}/Supplementary_Data.docx"); Q = S.paragraphs
assert Q[74].text.startswith("Figure S2. Plate position") and Q[75].text == "" and Q[76].text.startswith("Table S10.")
assert Q[80].text.startswith("Figure S3. Dispersion")
replace_once(Q[80], "Figure S3. Dispersion", "Figure S4. Dispersion")
CAP = ("Figure S3. Phylum-level composition of the 16S microbiota by tissue (a) at each station, all fish pooled, and (b) by genotypic "
       "category at the three stations where C. nasus, hybrids and P. toxostoma co-occur. Stacked bars give mean relative abundances of the "
       "twelve most abundant bacterial phyla (SILVA 138.2 nomenclature); remaining phyla are pooled as \"Other phyla\", and reads not assigned "
       "at phylum rank are shown as \"Unassigned\". Libraries with ≥ 3,000 reads after decontamination (1,784 libraries) were converted to "
       "relative abundances and averaged per individual × tissue across sequencing runs (628 profiles from 180 individuals). Numbers above "
       "bars are individuals. (a) All individuals of each station, genotypic categories pooled; each fish of a station weighs equally. "
       "Stations are ordered as in Fig. 1a; because their genotypic composition differs (Fig. 1a), a station bar combines station and "
       "category. Pseudomonadota is the most abundant phylum in 33 of the 36 tissue × station profiles, Fusobacteriota in the remaining "
       "three (hindgut at Saint-Just-d'Ardèche, midgut at Rosières and Saint-Just-d'Ardèche). Fusobacteriota and Bacillota are more "
       "abundant in both gut sections than in both external tissues at eight of nine stations, the exceptions being the Largue canal for "
       "Fusobacteriota and Avignon for Bacillota. The caudal fin bar of the Largue canal rests on three individuals. (b) Cn, C. nasus; Hy, "
       "all hybrids, intermediate and near-parental pooled; Pt, P. toxostoma. Stations: Canal (usine du Largue), Confluence Buëch-Méouge "
       "and Saint-Just-d'Ardèche. Because station is the main source of variation in community composition, categories are compared at "
       "equal stations: profiles were averaged per station and category, and each bar is the unweighted mean of the station means, so "
       "every station contributes equally and identically to the three categories. Within each tissue a station was kept only if every "
       "category had at least two individual profiles; for the caudal fin the Largue canal did not meet this criterion and the bars rest "
       "on two stations. Weighting by individuals instead of stations changes any bar segment by at most 6.5 percentage points. "
       "Intermediate and near-parental hybrids are pooled because they are unevenly distributed among stations (8 intermediate and 2 "
       "near-parental at Confluence Buëch-Méouge, 1 and 5 at Saint-Just-d'Ardèche): under the same two-profile criterion a four-category "
       "comparison would retain a single station per tissue. The unassigned fraction in the midgut bars of the Largue canal (a) and of "
       "C. nasus (b) comes from a single C. nasus individual in which one amplicon sequence variant unassigned at phylum rank accounts for "
       "80 % of midgut reads; it weighs one ninth of each of these bars. Correspondence with former names: Pseudomonadota = Proteobacteria; "
       "Bacillota = Firmicutes (including Mollicutes/Mycoplasmatales, formerly Tenericutes); Bacteroidota = Bacteroidetes; Fusobacteriota = "
       "Fusobacteria; Thermodesulfobacteriota includes the former Deltaproteobacteria (Desulfuromonadia, Desulfovibrionia, Desulfobulbia); "
       "Actinomycetota = Actinobacteria; Cyanobacteriota = Cyanobacteria. The figure is descriptive; contrasts between categories in "
       "community composition are tested in the main text (Fig. 3).")
anchor = Q[75]._p
blank1 = copy.deepcopy(Q[75]._p); img = copy.deepcopy(Q[75]._p); cap = copy.deepcopy(Q[74]._p)
for r in img.xpath("./w:r"): img.remove(r)
runs = cap.xpath("./w:r")
for r in runs[1:]: cap.remove(r)
anchor.addnext(blank1); anchor.addnext(cap); anchor.addnext(img)       # ordre final : 75, img, cap, blank1, Table S10
Paragraph(img, Q[75]._parent).add_run().add_picture(PNG, width=Inches(6.0))
cp = Paragraph(cap, Q[75]._parent); cp.runs[0].text = CAP
S2 = S.paragraphs
caps = [p.text[:10] for p in S2 if re.match(r"Figure S\d+\. ", p.text)]
assert caps == ["Figure S1.", "Figure S2.", "Figure S3.", "Figure S4."], caps
k = [i for i, p in enumerate(S2) if p.text.startswith("Figure S3. Phylum")][0]
assert S2[k - 1]._p.xpath(".//w:drawing") and S2[k + 2].text.startswith("Table S10."), "position de la figure"
assert S2[k].runs[0].italic and len(S2[k].runs) == 1
S.save(f"{OUTD}/Supplementary_Data.docx")
for f in MD5:
    print(f, hashlib.md5(open(f"{OUTD}/{f}", "rb").read()).hexdigest())
print("ordre des figures S dans l'article :", fig_order(TXT), "| légendes du supplément :", caps)
