#!/usr/bin/env python3
"""63-reponses_andre_1006b.py — second lot de réponses d'André du 2026-10-06 (transmises par JF) :
  - X = juvéniles (confirmé) ; X et cases vides -> NA dans la Table S2 (déjà fait par scripts/57) ; les fichiers de
    métadonnées du projet gardent X et vide ;
  - aucun numéro d'autorisation pour le Suran ni l'Ardèche : la mention de la Fédération de l'Ain et de l'ONEMA 07
    suffit (texte déjà en place, seules les notes grises changent) ;
  - Benjamin Hérodet nommé dans les remerciements (son rôle n'est pas précisé : laissé en note grise) ;
  - remerciements (reste) et contributions : vus en fin de manuscrit.
python-docx requis (absent sur meso) : exécuté hors cluster, fichiers déposés avec garde md5.
Usage : 63-reponses_andre_1006b.py <dossier_entree> <dossier_sortie>
"""
import sys, copy, hashlib
import docx
from docx.text.run import Run

IN, OUTD = sys.argv[1], sys.argv[2]
MD5 = {"Article.docx": "f211ce82c02641b0cf919ed93467115d", "Supplementary_Data.docx": "4b13d50bb29f6ac7babb94ec8bf85a1f"}
for f, h in MD5.items():
    assert hashlib.md5(open(f"{IN}/{f}", "rb").read()).hexdigest() == h, f"{f} : md5 inattendu"


def replace_once(par, old, new):
    hits = [r for r in par.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (old[:50], len(hits))
    hits[0].text = hits[0].text.replace(old, new)

OPEN = ("Reste ouvert : aucun numéro d'autorisation n'est donné pour le Suran ni pour l'Ardèche ; place de Benjamin Hérodet "
        "(Fédération de l'Ain) à préciser — déclaration d'éthique, remerciements ou auteurs.")
CLOSED = ("Suran et Ardèche : pas de numéro d'autorisation, la mention de la Fédération de l'Ain et de l'ONEMA 07 suffit ; "
          "Benjamin Hérodet (Fédération de l'Ain) est nommé dans les remerciements (André, 2026-10-06, second lot).")
A = docx.Document(f"{IN}/Article.docx"); P = A.paragraphs
assert P[105].text.strip() == "Ethics approval and consent to participate" and P[118].text.strip() == "Acknowledgements"
replace_once(P[19], OPEN, CLOSED)
replace_once(P[106], OPEN, CLOSED)
replace_once(P[115], "Section Acknowledgements : à rédiger (André) ; y nommer Benjamin Hérodet si c'est la place retenue.",
             "Benjamin Hérodet est nommé dans les remerciements (André, 2026-10-06) ; le reste des remerciements et les contributions "
             "des auteurs seront rédigés en fin de manuscrit.")
# remerciements : phrase en texte courant (mise en forme du paragraphe Funding), note grise conservée et mise à jour
ack = P[119]
assert ack.text == "[À rédiger en fin de manuscrit (André, 2026-10-03).]" and len(ack.runs) == 1
body = copy.deepcopy(P[114].runs[0]._r)
ack.runs[0]._r.addprevious(body)
Run(body, ack).text = "We thank Benjamin Hérodet (Fédération de l'Ain pour la pêche et la protection des milieux aquatiques). "
replace_once(ack, "[À rédiger en fin de manuscrit (André, 2026-10-03).]",
             "[Rôle de Benjamin Hérodet à préciser (aide aux pêches du Suran ?) ; autres remerciements à rédiger en fin de manuscrit "
             "(André, 2026-10-03 et 2026-10-06).]")
assert not ack.runs[0].italic and ack.runs[1].italic
txt = "\n".join(p.text for p in A.paragraphs)
assert "Reste ouvert : aucun numéro" not in txt and txt.count("Hérodet") == 5   # notes 19, 106, 115 ; remerciements : phrase + note, txt.count("Hérodet")
A.save(f"{OUTD}/Article.docx")

S = docx.Document(f"{IN}/Supplementary_Data.docx"); Q = S.paragraphs
assert Q[11].text.endswith("Sex: F, female; M, male; NA, sex not determined (juveniles, or sex not assigned).")
replace_once(Q[13], "correspondance [À CONFIRMER] par André.",
             "X = juvéniles confirmé par André (2026-10-06, second lot), qui demande NA pour X comme pour les cases vides dans la "
             "table et le maintien de X et du vide dans les fichiers de métadonnées.")
assert "[À CONFIRMER]" not in Q[13].text
S.save(f"{OUTD}/Supplementary_Data.docx")
for f in MD5:
    print(f, hashlib.md5(open(f"{OUTD}/{f}", "rb").read()).hexdigest())
