#!/usr/bin/env python3
"""73-additional_files.py — un Additional file par élément du supplément, pour Animal Microbiome.

Usage : 73-additional_files.py <Supplementary_Data.docx> <correspondance_additional_files.tsv> <ENA_accessions_durance.tsv>
        <sample_accessions.tsv> <ena_etat_final_oct.tsv> <dossier_figures> <dossier_sortie>

- tables -> Additional_file_n.xlsx : titre « Additional file n. … », paragraphes de légende à leur place (avant/après la
  table), en-tête en gras ; nombres écrits en nombres (sauf codes à zéro initial) ; notes grises retirées ; renvois S
  convertis, terminologie et italiques (texte riche) comme dans la copie de soumission ;
  Additional file 10 (Table S7) : les 2 304 runs du fichier préparé pour l'ENA à la place de l'extrait, plus l'accession
  d'échantillon (reçu ENA) et l'année de collecte de l'état ENA actuel (corrige 15Per2015Ch03A, pêché en 2014) ;
- notes -> Additional_file_n.docx : le bloc de la note dans une copie du supplément (styles conservés), titre renuméroté ;
- figures -> Additional_file_n.pdf : PDF vectoriels quand ils existent (S3, S4), sinon PNG -> PDF à 300 dpi (S1, S2, S5).
"""
import sys, os, re, csv, copy, shutil, hashlib
import docx, pandas as pd
from docx.text.paragraph import Paragraph
from docx.table import Table
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import soumission_commun as SC

SUP, CORF, ENAF, SAF, EFF, FIGD, OUT = sys.argv[1:8]
os.makedirs(OUT, exist_ok=True)
COR = {r["element"]: r for r in csv.DictReader(open(CORF), delimiter="\t")}
S = docx.Document(SUP)
body, ELEMS = SC.elements_supplement(S)
FIG = {"F1": ("fig_S1_depth_tradeoff.png", "png"), "F2": ("fig_S2_position_effect.png", "png"), "F3": ("fig_S3_phyla.pdf", "pdf"),
       "F4": ("fig_S4_ordination.pdf", "pdf"), "F5": ("fig_S5_permdisp.png", "png")}


def propre(t):
    return SC.terminologie(SC.conv(t))


def texte_par(el, lab=None):
    p = Paragraph(el, S)
    t = "".join(r.text for r in p.runs if not SC.est_note(r)).strip()
    return propre(SC.texte_T7(t) if lab == "T7" else t)


def riche(t, bold=False):
    segs = SC.segments_italiques(t)
    if not any(it for _, it in segs) and not bold: return t
    return CellRichText([TextBlock(InlineFont(i=it, b=bold), x) for x, it in segs])


def valeur(t):
    t = "" if t is None else str(t).strip()
    if t == "": return None
    if re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", t): return float(t) if "." in t else int(t)
    if re.fullmatch(r"[1-9]\d{0,2}(,\d{3})+", t): return int(t.replace(",", ""))
    return propre(t)


def table_T7():
    ena = pd.read_csv(ENAF, sep="\t", dtype=str)
    sa = pd.read_csv(SAF, sep="\t", dtype=str).apply(lambda s: s.str.strip())
    ef = pd.read_csv(EFF, sep="\t", dtype=str).apply(lambda s: s.str.strip())
    assert len(ena) == 2304 and sa.alias.is_unique
    m = ena.merge(sa.rename(columns={"alias": "sample_alias", "accession": "sample_accession"}), on="sample_alias",
                  how="left", validate="many_to_one")
    assert m.sample_accession.notna().all()
    cd = ef[ef.champ == "collection date"].set_index("accession").valeur
    an = m.sample_accession.map(cd).str[:4]
    ctrl = m.collection_date.str.startswith("missing")
    assert an[~ctrl].notna().all() and an[ctrl].isna().all()
    chg = (~ctrl) & (an != m.collection_date)
    assert set(m.loc[chg, "sample_alias"]) == {"DURANCE16S_15Per2015Ch03A"}, set(m.loc[chg, "sample_alias"])
    m.loc[~ctrl, "collection_date"] = an[~ctrl]
    cols = list(ena.columns); cols.insert(cols.index("sample_alias") + 1, "sample_accession")
    return m[cols].fillna(""), int(chg.sum())


def xlsx(lab, k0, k1, n, titre):
    wb = Workbook(); ws = wb.active; ws.title = f"Additional file {n}"
    ws.append([riche(f"Additional file {n}. {propre(titre)}.", bold=True)]); ws.append([])
    largeurs = {}; nt = 0
    for k in range(k0 + 1, k1):
        el = body[k]
        if el.tag.endswith("}p"):
            t = texte_par(el, lab)
            if t: ws.append([riche(t)]); ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=False)
            continue
        if not el.tag.endswith("}tbl"): continue
        nt += 1; tb = Table(el, S); rows = [[c.text for c in r.cells] for r in tb.rows]
        if lab == "T7":
            full, nchg = table_T7(); ex = pd.DataFrame(rows[1:], columns=rows[0]).rename(columns={"seq_replicate": "sequencing_replicate"})
            jo = ex.merge(full, on=list(ex.columns), how="left", indicator=True)
            assert (jo._merge == "both").all() and len(ex) == 3, "extrait absent du fichier complet"
            rows = [list(full.columns)] + full.values.tolist()
        ws.append([]) if ws.max_row > 2 else None
        ws.append([propre(h) for h in rows[0]])
        for c in ws[ws.max_row]: c.font = Font(bold=True)
        ws.freeze_panes = None
        for r in rows[1:]: ws.append([valeur(x) for x in r])
        for j, h in enumerate(rows[0]):
            largeurs[j] = max(largeurs.get(j, 0), min(60, max(len(str(x[j])) for x in rows) + 2))
        ws.append([])
    assert nt == 1, (lab, nt)
    from openpyxl.utils import get_column_letter
    for j, w in largeurs.items(): ws.column_dimensions[get_column_letter(j + 1)].width = w
    f = os.path.join(OUT, f"Additional_file_{n}.xlsx"); wb.save(f)
    txt = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert not SC.residuel(txt) and not SC.seize_s_isole(txt) and "excerpt" not in txt, (lab, SC.residuel(txt)[:3], SC.seize_s_isole(txt)[:3])
    return f, len(rows) - 1, len(rows[0])


def note(lab, k0, k1, n):
    D = docx.Document(SUP); b = list(D.element.body.iterchildren())
    for k, el in enumerate(b):
        if el.tag.endswith("}sectPr"): continue
        if not (k0 <= k < k1): el.getparent().remove(el)
    P = D.paragraphs
    for p in P:
        notes = [r for r in p.runs if SC.est_note(r)]
        for r in notes: r._r.getparent().remove(r._r)
        if notes and not p.text.strip(): p._p.getparent().remove(p._p)
    P = D.paragraphs
    h = P[0]; m = re.match(r"Note S\d+\. ", h.text); assert m, h.text
    pleins = [r for r in h.runs if r.text]
    assert len({(bool(r.italic), bool(r.bold)) for r in pleins}) == 1, [r.text for r in h.runs]
    nouveau = h.text.replace(m.group(0), f"Additional file {n}. ", 1)
    for r in h.runs:
        if r._r is not pleins[0]._r: r._r.getparent().remove(r._r)
    pleins[0].text = nouveau
    assert h.text.startswith(f"Additional file {n}. ")
    for p in D.paragraphs:
        SC.conv_paragraphe(p); SC.historique_paragraphe(p)
        for r in p.runs:
            r.text = SC.terminologie(r.text)
            if SC.couleur(r): r.font.color.rgb = None
        for r in list(p.runs): SC.italiciser_run(p, r)
    for t in D.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    SC.conv_paragraphe(p)
                    for r in p.runs: r.text = SC.terminologie(r.text)
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    xml = D.element.xml
    for rid, rel in list(D.part.rels.items()):
        if rel.reltype == RT.IMAGE and f'"{rid}"' not in xml: D.part.drop_rel(rid)
    assert not D.element.xpath(".//w:drawing")
    assert not [r for r in D.part.rels.values() if r.reltype == RT.IMAGE]
    f = os.path.join(OUT, f"Additional_file_{n}.docx"); D.save(f)
    txt = "\n".join(p.text for p in docx.Document(f).paragraphs)
    assert not SC.residuel(txt), SC.residuel(txt)[:5]
    assert not SC.seize_s_isole(txt), SC.seize_s_isole(txt)[:5]
    return f, len(docx.Document(f).paragraphs)


def figure(lab, n):
    nom, typ = FIG[lab]; src = os.path.join(FIGD, nom); f = os.path.join(OUT, f"Additional_file_{n}.pdf")
    if typ == "pdf": shutil.copyfile(src, f)
    else:
        im = Image.open(src); im = im.convert("RGB") if im.mode != "RGB" else im
        im.save(f, "PDF", resolution=300.0)
    return f, nom


bilan = []
for k0, k1, lab, titre in ELEMS:
    n = SC.AF[lab]; assert COR[lab]["additional_file"] == str(n)
    if lab[0] == "T": f, nr, nc = xlsx(lab, k0, k1, n, titre); info = f"{nr} lignes x {nc} colonnes"
    elif lab[0] == "N": f, npar = note(lab, k0, k1, n); info = f"{npar} paragraphes"
    else: f, src = figure(lab, n); info = f"source {src}"
    md5 = hashlib.md5(open(f, "rb").read()).hexdigest(); taille = os.path.getsize(f)
    assert taille < 20e6, (f, taille)
    bilan.append((n, lab, os.path.basename(f), taille, md5, info))
bilan.sort()
assert all(v == 1 for v in SC.HIST_COMPTE.values()), SC.HIST_COMPTE
with open(os.path.join(OUT, "bilan_additional_files.tsv"), "w") as fh:
    fh.write("additional_file\telement\tfichier\toctets\tmd5\tcontenu\n")
    for b in bilan: fh.write("\t".join(map(str, b)) + "\n")
for b in bilan: print(*b, sep=" | ")
