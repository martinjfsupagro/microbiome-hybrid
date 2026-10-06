"""57-sexe_NA_tableS2.py — demande d'André (visio du 2026-10-06, transmise par JF) : « mettre NA partout » pour
les individus non sexés. Table S2 du supplément : colonne Sex, X (52) et – (17) -> NA (69) ; légende et note grise
mises à jour. Les autres colonnes gardent « – » (non applicable ou non disponible).

Hors champ, délibérément : les métadonnées du projet (metadata/*.csv) gardent X et vide, qui distinguent encore
juvéniles et adultes non sexés (taille moyenne 13,0 cm contre 18,2 cm) ; aucun attribut de sexe n'est déposé à
l'ENA. Seul scripts/29 lit le sexe, et seulement F et M.
python-docx absent de meso : exécution hors cluster. Garde md5 en entrée.
"""
import hashlib, os, collections
import docx

M = os.environ.get("MS_DIR", "docs/manuscrit/"); OUT = os.environ.get("OUT_DIR", M)
F = "Supplementary_Data.docx"; MD5_IN = "1b677bb0cb7b601adaa49e74f80708c3"     # commit 2a46d4c
assert hashlib.md5(open(M + F, "rb").read()).hexdigest() == MD5_IN, "md5 d'entrée inattendu"
s = docx.Document(M + F)
T = [t for t in s.tables if [c.text for c in t.rows[0].cells][15:16] == ["Sex"]]; assert len(T) == 1
avant = collections.Counter(r.cells[15].text for r in T[0].rows[1:])
assert avant == {"M": 78, "X": 52, "F": 33, "–": 17}, avant
for r in T[0].rows[1:]:
    c = r.cells[15]
    if c.text in ("X", "–"):
        runs = c.paragraphs[0].runs; assert len(c.paragraphs) == 1 and len(runs) == 1
        runs[0].text = "NA"
apres = collections.Counter(r.cells[15].text for r in T[0].rows[1:])
assert apres == {"M": 78, "NA": 69, "F": 33}, apres

leg = [p for p in s.paragraphs if p.text.startswith("Genotypic category of each of the 180")]; assert len(leg) == 1
R = [r for r in leg[0].runs if "X and –, sex not determined" in r.text]; assert len(R) == 1
R[0].text = R[0].text.replace("X and –, sex not determined (juveniles, or sex not assigned).",
                              "NA, sex not determined (juveniles, or sex not assigned).")
G = [r for p in s.paragraphs for r in p.runs if r.text.lstrip().startswith("[Codes de sexe :")]; assert len(G) == 1
G[0].text = ("[Codes de sexe : les individus non sexés sont notés NA (69), à la demande d'André (2026-10-06) ; les "
             "métadonnées du projet gardent X (52, taille moyenne 13,0 cm, juvéniles présumés) et vide (17, 18,2 cm, adultes "
             "non sexés présumés) — correspondance [À CONFIRMER] par André. Table insérée le 2026-10-03 sur décision de JF "
             "(question 2.9) ; les tables suivantes sont renumérotées S3 à S10.]")
s.save(OUT + F)
print("Table S2, colonne Sex :", dict(avant), "->", dict(apres))
print("écrit :", OUT + F, hashlib.md5(open(OUT + F, "rb").read()).hexdigest())
