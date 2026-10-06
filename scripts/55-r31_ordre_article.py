"""55-r31_ordre_article.py — intégration au manuscrit du test R31 (contrôle v) et de l'effet de l'ordre de
traitement sur l'alpha du hindgut, décision de JF du 2026-10-06.

Source des chiffres : results/tests_20261006/alpha_ordre/alpha_ordre.tsv (scripts/51, plan
docs/plan_R31_ordre_2026-10-06.md, commit f3916c1). Les chaînes insérées sont construites depuis cette
table et comparées aux valeurs attendues (assertions) ; scripts/40 les recontrôle.

  §8.12      contrôle (v) décrit, avec sa chronologie réelle : déclaré après les résultats de (iii),
             avant son propre calcul.
  Results    (Alpha diversity) : verdict du contrôle (v) ; baisse de l'alpha avec le rang dans le hindgut
             et, sur les indices de présence, dans le midgut ; absence de biais pour les contrastes du
             hindgut ; note grise remplacée.
  Discussion (limites) : la phrase « which was not tested » est remplacée. Corrige au passage l'italique
             de « C. nasus », perdu le matin même par le remplacement à travers les runs du script 49
             (erreur de l'agent).

Les noms d'espèce sont insérés comme runs italiques distincts ; tout texte nouveau est bleu.
python-docx absent de meso : exécution hors cluster. Garde md5 en entrée.
"""
import copy, hashlib, os
import numpy as np, pandas as pd
import docx
from docx.shared import RGBColor
from docx.text.run import Run

M = os.environ.get("MS_DIR", "docs/manuscrit/"); OUT = os.environ.get("OUT_DIR", M)
TAB = os.environ.get("R31_TSV", "results/tests_20261006/alpha_ordre/alpha_ordre.tsv")
MD5_IN = "b4ec479023fa85bf58a2c04e4af7de86"                       # commit 2a46d4c
BLUE, GREY = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x88, 0x88, 0x88)

# ── chiffres ────────────────────────────────────────────────────────────────────────────────
R = pd.read_csv(TAB, sep="\t", dtype={"passe": str})
est = R[(R.classe_principal == "dominant Cn") & (((R.passe == "1") & (R.tissu == "caudale")) | ((R.passe == "2") & (R.tissu == "midgut")))].copy()
assert len(est) == 7, len(est)
est["ratio"] = (est.delta_a_Hy_moins_Pt / est.delta_r_Hy_moins_Pt).abs() / np.maximum(est.beta_w_bas.abs(), est.beta_w_haut.abs())
assert ((est.delta_a_Hy_moins_Pt / est.delta_r_Hy_moins_Pt < est.beta_w_bas) | (est.delta_a_Hy_moins_Pt / est.delta_r_Hy_moins_Pt > est.beta_w_haut)).all()
dr_c = abs(R[(R.passe == "1") & (R.tissu == "caudale")].delta_r_Hy_moins_Pt.iloc[0])
dr_m = abs(R[(R.passe == "2") & (R.tissu == "midgut")].delta_r_Hy_moins_Pt.iloc[0])
H = R[R.tissu == "hindgut"]; assert len(H) == 8 and (H.p_beta_w < 0.05).all() and (H.beta_w < 0).all()
pct = (1 - np.exp(10 * H.beta_w)) * 100
Mg = R[(R.tissu == "midgut") & (R.p_beta_w < 0.05)]
assert len(Mg) == 3 and (~Mg.ponderee.astype(str).isin(["True"])).all() and (Mg.beta_w < 0).all()
dr_h = sorted(abs(x) for x in H.groupby("passe").delta_r_Hy_moins_Pt.first())
V = dict(drc=f"{dr_c:.1f}", drm=f"{dr_m:.1f}", rmin=f"{est.ratio.min():.0f}", rmax=f"{est.ratio.max():.0f}",
         pmax=f"{H.p_beta_w.max():.3f}", pmin_=f"{pct.min():.0f}", pmax_=f"{pct.max():.0f}",
         drh1=f"{dr_h[1]:.1f}", drh0=f"{dr_h[0]:.1f}")
assert V == dict(drc="0.5", drm="0.4", rmin="16", rmax="37", pmax="0.008", pmin_="25", pmax_="46", drh1="0.2", drh0="0.1"), V


# ── outils ──────────────────────────────────────────────────────────────────────────────────
def para(doc, start):
    P = [p for p in doc.paragraphs if p.text.startswith(start)]
    assert len(P) == 1, (start, len(P)); return P[0]


def insert_after(p, anchor, segments):
    """Insère après le run `anchor` des runs bleus ; segments = [(texte, italique)]."""
    tmpl = [r for r in p.runs if not r.italic and r.text][0]
    prev = anchor._r
    for text, it in segments:
        n = copy.deepcopy(tmpl._r); prev.addnext(n); prev = n
        r = Run(n, p); r.text = text; r.italic = True if it else None; r.font.color.rgb = BLUE


def segs(s):
    """'… <i>P. toxostoma</i> …' -> [(texte, italique)]"""
    out = []
    for i, part in enumerate(s.replace("</i>", "<i>").split("<i>")):
        if part: out.append((part, i % 2 == 1))
    return out


assert hashlib.md5(open(M + "Article.docx", "rb").read()).hexdigest() == MD5_IN, "md5 d'entrée inattendu"
d = docx.Document(M + "Article.docx")

# ── §8.12 ───────────────────────────────────────────────────────────────────────────────────
p = para(d, "Four further controls were declared")
last = p.runs[-1]; assert last.text.endswith("within each fishing campaign."), last.text[-60:]
insert_after(p, last, segs(
    " A fifth control was declared in the same way, in a document committed after the results of (iii) were "
    "known but before its own computation. (v) Whether processing order could produce the classification of "
    "hybrid alpha diversity: fish were ranked by individual number within each campaign (1 = first dissected; "
    "the two 2014 Pertuis series, fished on different days, were ranked separately); the slope of log alpha "
    "diversity on rank was estimated within station × category cells, where it cannot carry a category effect, "
    "and compared with the slope that rank alone would require to produce the hybrid − <i>P. toxostoma</i> "
    "contrast, given the mean rank difference between hybrids and <i>P. toxostoma</i> within stations."))

# ── Results, alpha ──────────────────────────────────────────────────────────────────────────
p = para(d, "Hybrids were transgressive in none of the 32 combinations")
anchor = [r for r in p.runs if r.text == ", not an equality."]; assert len(anchor) == 1
insert_after(p, anchor[0], segs(
    f" Processing order could not produce either classification (§8.12, control v): within stations, hybrids "
    f"and <i>P. toxostoma</i> were dissected at the same mean rank (differences of {V['drc']} ranks in the caudal "
    f"fin in pass 1 and {V['drm']} in the midgut in pass 2), so that the slope required to produce the hybrid − "
    f"<i>P. toxostoma</i> contrast exceeded the within-cell slope on rank by a factor of {V['rmin']} to "
    f"{V['rmax']} on every index. Alpha diversity nevertheless declined with processing rank within station × "
    f"category cells in the hindgut, on all four indices in both passes (p ≤ {V['pmax']}), by {V['pmin_']}–"
    f"{V['pmax_']} % per ten ranks, and in the midgut on presence-based indices only (3 of 8 tests). Hybrids and "
    f"<i>P. toxostoma</i> having been processed at the same mean rank in the hindgut (differences of "
    f"{V['drh1']} and {V['drh0']} ranks), this decline does not bias the hindgut category contrasts; the time "
    f"spent in the holding tank, which grows with rank, is consistent with it but cannot be separated from plate "
    f"column in this design."))
G = [r for r in p.runs if r.text.lstrip().startswith("[Ordre de traitement testé")]; assert len(G) == 1
G[0].text = (" [Contrôle (v) : plan docs/plan_R31_ordre_2026-10-06.md (f3916c1), scripts/51, R36 ; intégré sur décision "
             "de JF du 2026-10-06. Contraste femelles/mâles (R27) non rapporté : décision de JF attendue — les femelles "
             "ont été disséquées plus tôt (rang médian 5 contre 8, p = 0,02), et l'ordre abaisse l'alpha digestive. "
             "À retirer avant soumission.]")

# ── Discussion, limites (corrige aussi l'italique perdu par le script 49) ───────────────────
p = para(d, "Four limits bound these conclusions.")
r0, r1, r2, r3, r4 = p.runs[:5]
OLD0 = "Because C. nasus was generally processed first, processing order could contribute to"
assert r0.text.endswith(OLD0) and r1.text == "" and r2.text == " the dominance of " and r3.text == "C. nasus"
OLD4 = " in caudal-fin alpha diversity, which was not tested. "
assert r4.text.startswith(OLD4 + "Rarefaction"), r4.text[:80]
r0.text = r0.text[: -len(OLD0)]
r4.text = r4.text[len(OLD4):]
for r in (r1, r2, r3): r._r.getparent().remove(r._r)
insert_after(p, r0, segs(
    "Hindgut alpha diversity declined with processing rank within station and category, which is consistent with "
    "this proposal without testing it, rank being confounded with column; processing order could not, however, produce the dominance "
    "of <i>C. nasus</i> in alpha diversity, although <i>C. nasus</i> was generally processed first, hybrids and "
    "<i>P. toxostoma</i> having been dissected at the same mean rank within stations. "))

# ── contrôle : aucun nom d'espèce hors italique ──────────────────────────────────────────────
import re
bad = [(k, r.text) for k, q in enumerate(d.paragraphs) for r in q.runs
       if re.search(r"\b(C\. nasus|P\. toxostoma)\b", r.text) and not r.italic]
assert not bad, bad
d.save(OUT + "Article.docx")
print("écrit :", OUT + "Article.docx", hashlib.md5(open(OUT + "Article.docx", "rb").read()).hexdigest())
print(V)
