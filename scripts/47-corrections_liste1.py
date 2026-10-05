"""47-corrections_liste1.py — corrections du 2026-10-05 (soir), sur les deux versions de l'article et du supplément.
1. §5 (et Note S2 de la version resserrée) : part du 12S de l'hôte et des dimères d'amorces remplacée par les valeurs
   de scripts/45-resume_12S.py (results/verif_article/part_12S.tsv) ; « 15–67 % » n'était reproduit par aucun calcul.
2. Note grise de l'introduction réduite à ce qui reste ouvert.
3. Supplément : note grise de la Figure S1 corrigée (« ex-Figure S2 » ; le renumérotage du script 44 l'avait inversée) ;
   largeur des images ramenée à la zone de texte (15,24 cm), rapport hauteur/largeur conservé."""
import docx, hashlib, os, csv, re
from docx.shared import RGBColor, Pt
from docx.oxml.ns import qn
M = os.environ.get("MS_DIR", "docs/manuscrit/"); OUT = os.environ.get("OUT_DIR", M)
P12 = os.environ.get("P12", "results/verif_article/part_12S.tsv")
IN = {"Article.docx": "7b0cdc88e112dab2d7d63c0e156053d1", "Article_resserre.docx": "1b0a55c5cf2f951427f9672f807564d7",
      "Supplementary_Data.docx": "49487daa7c84222243c11b44d6a1a544", "Supplementary_Data_resserre.docx": "9ebff25d57171574125f64a1fdd676f5"}
for f, h in IN.items():
    assert hashlib.md5(open(M + f, "rb").read()).hexdigest() == h, f"md5 inattendu : {f}"
BLUE = RGBColor(0x1F, 0x4E, 0x79); GREY = RGBColor(0x88, 0x88, 0x88)
R = list(csv.DictReader(open(P12), delimiter="\t"))
runs = [r for r in R if r["run"] != "tous"]; tous = [r for r in R if r["run"] == "tous"][0]
g = [float(r["pct_global"]) for r in runs]
NEW12 = (f"(host 12S and primer dimers: {min(g):.0f}–{max(g):.0f} % of reads per run; median {float(tous['mediane']):.0f} % "
         f"per biological sample, range {float(tous['minimum']):.0f}–{float(tous['maximum']):.0f} %)")
OLD12 = "(15–67% of reads per sample)"
INTRO = ("[Introduction — reste ouvert (note réduite le 2026-10-05 ; points réglés retirés : structure Animal Microbiome, "
         "description du 4H au §8.11, effectifs des quasi-purs, corrections Small 2019 et Sevellec 2019 propagées dans la .bib). "
         "(i) Titre provisoire (sous-titre du supplément, JF 2026-10-03) ; titre définitif en fin d'écriture. "
         "(ii) Wang et al. 2015 [11] et Sevellec et al. 2014 [6] viennent de la liste du manuscrit d'origine et n'ont pas été "
         "relus en texte intégral : vérifier leur pertinence là où ils sont cités.]")
MAXW = 5486400  # 15,24 cm en EMU
LOG = []

def isgrey(r): return r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) == "888888"

def blue_sub(p, old, new):
    hit = [r for r in p.runs if old in r.text]; assert len(hit) == 1, (old, len(hit))
    r = hit[0]; a, b = r.text.split(old, 1); r.text = a
    import copy
    from docx.text.run import Run
    n1 = copy.deepcopy(r._r); r._r.addnext(n1); n2 = copy.deepcopy(r._r); n1.addnext(n2)
    R1, R2 = Run(n1, p), Run(n2, p); R1.text = new; R1.font.color.rgb = BLUE; R2.text = b

for f in IN:
    d = docx.Document(M + f)
    hits = [p for p in d.paragraphs if OLD12 in p.text]
    for p in hits: blue_sub(p, OLD12, NEW12)
    LOG.append(f"{f}: 12S remplacé dans {len(hits)} paragraphe(s)")
    if f.startswith("Article"):
        assert len(hits) == 1, f
        p = [p for p in d.paragraphs if "[Introduction, draft 1" in p.text]; assert len(p) == 1, f; p = p[0]
        grey = [r for r in p.runs if isgrey(r)]; assert grey and all(isgrey(r) for r in p.runs if r.text.strip()), f
        for r in grey[1:]: r._r.getparent().remove(r._r)
        grey[0].text = INTRO
    else:
        p = [p for p in d.paragraphs if "(ex-Figure S1)" in p.text]; assert len(p) == 1, f
        for r in p[0].runs:
            if "(ex-Figure S1)" in r.text: r.text = r.text.replace("(ex-Figure S1)", "(ex-Figure S2)")
        n = 0
        for inl in d.element.body.iter(qn("wp:inline")):
            ext = inl.find(qn("wp:extent")); cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            if cx > MAXW:
                ncy = int(cy * MAXW / cx); ext.set("cx", str(MAXW)); ext.set("cy", str(ncy))
                for e in inl.iter(qn("a:ext")): e.set("cx", str(MAXW)); e.set("cy", str(ncy))
                n += 1
        LOG.append(f"{f}: note ex-Figure corrigée ; {n} image(s) ramenée(s) à 15,24 cm")
    d.save(OUT + f)
print(NEW12); print("\n".join(LOG))
