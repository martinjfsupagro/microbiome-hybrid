"""Génère points_a_regler_2026-10-07.docx (accessible : styles de titres, listes balisées,
en-têtes de tableaux répétés, métadonnées, langue fr-FR, Arial 12, interligne 1,5)."""
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "points_a_regler_2026-10-07.docx"
cr = pd.read_csv("crossref_years.csv")

doc = Document()
cp = doc.core_properties
cp.title = "microbiome-hybrid : points à régler avant la soumission à Animal Microbiome"
cp.subject = "Liste des décisions, informations et corrections en attente, état du 7 octobre 2026"
cp.language = "fr-FR"
cp.author = "Projet microbiome-hybrid"
cp.keywords = "microbiome-hybrid; soumission; références; Crossref"

def set_lang(style, lang="fr-FR"):
    rpr = style.element.get_or_add_rPr()
    l = rpr.find(qn("w:lang"))
    if l is None:
        l = OxmlElement("w:lang"); rpr.append(l)
    l.set(qn("w:val"), lang)

for name in ["Normal", "List Bullet", "List Number", "Title", "Heading 1", "Heading 2", "Heading 3"]:
    st = doc.styles[name]
    st.font.name = "Arial"
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    set_lang(st)
    pf = st.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if name in ("Normal", "List Bullet", "List Number"):
        st.font.size = Pt(12); pf.line_spacing = 1.5; pf.space_after = Pt(6)
for name, sz in [("Title", 20), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
    st = doc.styles[name]; st.font.size = Pt(sz); st.font.bold = True
    st.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)
    st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)

def P(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text); return p
def B(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix: p.add_run(bold_prefix).bold = True
    p.add_run(text); return p
def H(text, lvl): return doc.add_heading(text, level=lvl)

def table(header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, h in enumerate(header):
        c = hdr.cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "1565C0")
        c._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs: r.font.size = Pt(10); r.font.name = "Arial"
    # résumé du tableau pour les lecteurs d'écran
    tblPr = t._tbl.tblPr
    cap = OxmlElement("w:tblCaption"); cap.set(qn("w:val"), header[0] + " : " + ", ".join(header[1:])); tblPr.append(cap)
    doc.add_paragraph()
    return t

# ------------------------------------------------------------------ titre
doc.add_heading("microbiome-hybrid : points à régler avant la soumission", 0)
P("État au 7 octobre 2026. Sources : dossier _etat/ du dépôt sur meso (commit c9e6169, identique au miroir GitHub public), "
  "version de travail docs/manuscrit/Article.docx, paquet docs/soumission/, et une vérification des 51 références de "
  "l'article contre Crossref faite le 7 octobre (section 3).")
P("Chaque point indique ce qu'il faut faire, qui doit le faire et s'il bloque la soumission. Les points notés « à confirmer » "
  "sont des hypothèses, pas des constats.")

# ------------------------------------------------------------------ 1. vue d'ensemble
H("1 Vue d'ensemble", 1)
table(["N°", "Point", "Qui", "Bloque la soumission"], [
    ["2.1", "Huit informations administratives (auteurs, adresses, courriels, auteur correspondant, DOI Zenodo, rôle des financeurs, contributions, autres remerciements)", "JF", "Oui"],
    ["2.2", "Titre définitif", "JF", "Oui"],
    ["2.3", "Rôle de B. Hérodet dans la phrase de remerciement", "JF ou André", "Non"],
    ["3.1", "Convention d'année des références (volume imprimé) : à confirmer", "JF", "Oui (forme)"],
    ["3.2", "Sinama et al. 2013 [28] : quelle notice citer", "JF", "Oui (forme)"],
    ["3.3", "Pertinence de Wang 2015 [11] et Sevellec 2014 [6]", "JF", "Non"],
    ["3.4", "Défauts de forme ou mises à jour dans 13 références", "Agent, après accord de JF", "Oui (forme)"],
    ["4", "Appel à la Table S5 au §8.6 ; « Eukaryota » et « Mitochondria » en romain ; lettre d'accompagnement, relecteurs, résumé graphique", "JF", "Lettre : oui ; le reste : non"],
    ["5", "Part des contributions d'André", "André", "Oui"],
    ["6", "Adapter scripts/72, régénérer le paquet, relancer les contrôles, archiver sur Zenodo", "Agent", "Oui"],
    ["7", "ENA : corriger un TITLE, puis rendre le projet public", "Humain au clavier (identifiants Webin)", "Ouverture publique : oui"],
    ["8", "Points facultatifs (ASV0185, Mycoplasmatales, 12S, mécanismes, hygiène)", "—", "Non"],
])

# ------------------------------------------------------------------ 2
H("2 Informations à fournir", 1)
H("2.1 Les huit marqueurs [TO COMPLETE]", 2)
P("Le paquet de soumission (docs/soumission/Manuscript_microbiome-hybrid.docx) contient huit marqueurs à remplacer :")
for t in ["noms complets de tous les auteurs, avec numéros d'affiliation ;",
          "adresses institutionnelles ;",
          "adresses électroniques ;",
          "auteur correspondant (nom et courriel) ;",
          "DOI de la version archivée du code sur Zenodo (voir section 6) ;",
          "rôle des financeurs (EDF, contrat EDF-CNRS AGDI 428481 ; Fédération de l'Ain) dans la conception, la collecte, l'analyse, l'interprétation et la rédaction ;",
          "contribution de chaque auteur, par initiales ;",
          "autres remerciements (B. Hérodet y est déjà nommé)."]:
    B(t)
P("Ces marqueurs ne sont pas dans Article.docx : scripts/72 les insère en dur à chaque génération du paquet. Il suffit "
  "d'envoyer ces informations, en texte simple ; l'agent adapte le script pour qu'il les reprenne (section 6). Ne pas les "
  "saisir dans Article.docx : un texte de contributions saisi là fait échouer la génération du paquet.",
  "Important. ")
H("2.2 Titre définitif", 2)
P("Titre actuel, provisoire depuis le 3 octobre : « Host-associated bacterial microbiota of hybridising cyprinid fishes ».")
H("2.3 Benjamin Hérodet", 2)
P("Décision du 7 octobre : il figure dans les remerciements. La phrase actuelle est « We thank Benjamin Hérodet (Fédération "
  "de l'Ain pour la pêche et la protection des milieux aquatiques). » Son rôle n'est pas précisé ; « aide aux pêches du "
  "Suran » est une hypothèse à confirmer. Si personne ne le précise, la phrase reste telle quelle.")

# ------------------------------------------------------------------ 3 références
H("3 Références", 1)
H("3.1 Dates : vérification contre Crossref", 2)
P("Convention appliquée dans l'article (note grise des références) : l'année citée est celle du volume imprimé quand il "
  "existe, et non celle de la mise en ligne. La note grise demande de confirmer cette règle et nomme quatre références "
  "concernées ; il y en a en fait sept.")
P("Résultat de la vérification : sur les 48 références de l'article qui ont un DOI, l'année citée est l'année du volume "
  "imprimé ou, à défaut, l'année de publication enregistrée par Crossref, avec une seule exception (SILVA [36], ci-dessous). "
  "Les 48 numéros de volume sont conformes à Crossref, et les 45 paginations ou numéros d'article qu'il fournit aussi. "
  "Aucune date n'est donc à corriger si la règle est confirmée.")
lab = {21: "Nielsen et al., Ecol Lett 26", 24: "Čížková et al., Mol Ecol 33", 29: "Guivier et al., Hydrobiologia 830",
       30: "Seehausen et al., Mol Ecol 17", 31: "Guivier et al., Microb Ecol 75", 45: "Anderson, Austral Ecol 26",
       46: "Anderson, Biometrics 62"}
rows = []
for n, l in lab.items():
    r = cr[cr.n == n].iloc[0]
    note = "Mise en ligne rétrospective (numérisation)" if n == 45 else "Conforme à la règle"
    rows.append([f"[{n}]", l, int(r.y_online), int(r.y_print), int(r.cited_year), note])
table(["Réf.", "Référence", "En ligne", "Volume imprimé", "Année citée", "Verdict"], rows)
P("Je n'ai pas relu les consignes de la revue sur ce point (docs/soumission/consignes/). La convention du volume imprimé "
  "est la plus courante pour un article paru dans un volume paginé.", "Décision attendue de JF : confirmer la règle. ")
P("Crossref ne donne pour [36] qu'une mise en ligne en 2012 (année « issued » 2012), sans date d'impression. L'article cite "
  "2013, année du volume 41, numéro « Database issue ». C'est cohérent avec la règle, mais Crossref ne permet pas de le "
  "vérifier : contrôler sur la page de l'éditeur.", "SILVA, Quast et al. [36]. ")

H("3.2 Sinama et al. 2013 [28]", 2)
P("Notice citée : Sinama M, Gilles A, Costedoat C, Corse E, Olivier J-M, Chappaz R, Pech N. Non-homogeneous combination of "
  "two porous genomes induces complex body shape trajectories in cyprinid hybrids. Front Zool. 2013;10:22. L'année est "
  "conforme à Crossref (volume imprimé 2013).")
P("La question n'est donc pas la date mais le choix de la notice. Le 5 octobre, deux notices étaient possibles (la thèse "
  "de M. Sinama ou cet article de Frontiers in Zoology) ; l'article a été retenu, à confirmer. La phrase qui le cite :")
P("« Hybridisation occurs in the patchily distributed sympatric areas of the basin, with an incidence unrelated to habitat "
  "characteristics [27–29] — the configuration of a mosaic hybrid zone [30]. »")
P("L'article de Frontiers in Zoology porte sur la forme du corps des hybrides. S'il ne traite pas de l'incidence de "
  "l'hybridation selon l'habitat, la thèse est peut-être la bonne source pour cette phrase ; il faudrait alors donner sa "
  "notice complète (année, université).", "Décision attendue de JF : ")

H("3.3 Pertinence de deux références", 2)
B("« Introgressive hybridisation may disrupt co-adapted gene complexes and thereby alter close host–microbe interactions "
  "[11]; … » (Background). Référence : Wang et al., Nat Commun 2015 (souris hybrides).", "Wang 2015 [11]. ")
B("« … surveys of wild populations accordingly report a predominant environmental component of variation in bacterial "
  "composition [6–8]. » (Background). Référence : Sevellec et al., J Evol Biol 2014 (corégones). Sevellec et al. 2019 "
  "[19] est cité ailleurs dans sa version publiée.", "Sevellec 2014 [6]. ")
P("Décision attendue de JF : garder, remplacer ou retirer chacune. Retirer une référence oblige à renuméroter la liste.")

H("3.4 Défauts de forme trouvés le 7 octobre", 2)
P("Trouvés en contrôlant la liste. Aucun ne change le sens. Correction proposée dans Article.docx, par script et avec "
  "contrôle md5 de la version d'entrée, puis régénération du paquet.")
table(["Réf.", "Défaut", "Correction proposée"], [
    ["[3], [10], [17], [33], [43]", "Identifiant d'article mSystems, mBio ou mSphere écrit avec un tiret demi-cadratin (ex. e01785–15)", "Trait d'union, comme dans le DOI : e01785-15, e00331-19, e01181-22, e00032-16, e00354-23"],
    ["[27], [30]", "Noms d'auteurs en capitales (COSTEDOAT C, PECH N… ; SEEHAUSEN O…), repris tels quels de Crossref", "Costedoat C, Pech N, Salducci M-D, Chappaz R, Gilles A ; Seehausen O, Takimoto G, Roy D, Jokela J"],
    ["[5]", "Front Microbiol. 2014;5 : pas de numéro d'article (Crossref n'en fournit pas)", "2014;5:207 d'après le DOI (fmicb.2014.00207) ; à vérifier sur la page de l'éditeur"],
    ["[38]", "« et al.. » : point doublé", "« et al. »"],
    ["[47]", "vegan sans année", "Ajouter 2026 : la version 2.7-5 est publiée sur CRAN le 25/05/2026 (DESCRIPTION de l'env dada2 sur meso)"],
    ["[22]", "Donohue et al., prépublication bioRxiv (déposée le 11/12/2025) ; au 07/10/2026, Crossref ne la relie à aucune version publiée", "Revérifier juste avant la soumission ; citer la version publiée si elle existe"],
    ["[26]", "Thèse de Gilles 1998 : ni université ni ville", "Ajouter l'établissement ; à fournir par JF"],
    ["[51]", "Dépôt GitHub « Accessed 3 Oct 2026 »", "Mettre la date de soumission ; remplacer ou compléter par le DOI Zenodo (2.1)"],
])
P("HybridMicrobiomes [49] : version 0.1.1, publiée sur CRAN le 05/12/2023, citée 2023 : conforme.")

# ------------------------------------------------------------------ 4
H("4 Manuscrit : points de forme (JF)", 1)
B("au §8.6, la phrase « Monte-Carlo error … (Supplementary Table S5; details in Supplementary Note S2) » "
  "devient « Additional file 7 » dans le paquet : faut-il garder cet appel ?", "Table S5 : ")
B("ce sont les libellés de la taxonomie SILVA, laissés en romain : à confirmer.", "« Eukaryota » et « Mitochondria » : ")
B("lettre d'accompagnement (obligatoire), relecteurs suggérés, résumé graphique (facultatif).", "Pièces jointes : ")

# ------------------------------------------------------------------ 5
H("5 Attendu d'André", 1)
B("Sa part des contributions des auteurs (à intégrer à la réponse de 2.1).")

# ------------------------------------------------------------------ 6
H("6 Préalables techniques (agent)", 1)
for t in ["Adapter scripts/72 pour qu'il lise page de titre, contributions, financeurs, DOI et remerciements dans un fichier séparé, et ne garde de marqueur que pour un champ vide. Réduire au passage le marqueur Hérodet à « other acknowledgements ».",
          "Appliquer les corrections de la section 3.4 et les décisions des sections 2 à 4 dans Article.docx.",
          "Régénérer le paquet (scripts 72, 73, 75 ; poste local, python-docx absent de meso), puis relancer les contrôles du texte (scripts 40 et 30).",
          "Archiver le dépôt sur Zenodo (DOI à citer en [51] et dans « Availability of data and materials »), puis resynchroniser le miroir GitHub."]:
    doc.add_paragraph(t, style="List Number")

# ------------------------------------------------------------------ 7
H("7 Dépôt ENA PRJEB124417 (humain au clavier)", 1)
B("le TITLE de ERS31168490 (échantillon 15Per2015Ch03A, pêché le 20/08/2014) indique « Per 2015 » au lieu de « Per 2014 ». "
  "Faire un MODIFY de ce seul champ, construit depuis l'état déposé (méthode de scripts/53), en test puis en production. "
  "Corriger aussi la valeur de référence avant de revérifier (7_verifier_exhaustif.sh). Procédure : "
  "ena_deposit/MEMO_corrections_restantes.md.", "Correction : ")
B("rendre le projet public au moment de la soumission (décision du 03/10).", "Ouverture publique : ")

# ------------------------------------------------------------------ 8
H("8 Points facultatifs, non bloquants", 1)
for t in ["Identité de l'ASV0185 (phylum non assigné ; 80 % des lectures midgut de 2015_Caa_1003) : BLAST non fait.",
          "Part des Mycoplasmatales, comptées dans Bacillota par SILVA 138.2.",
          "Excès de 12S en plaque 2, colonnes 06–08 (R40) : mécanisme à confirmer.",
          "Mécanismes de R22, R29 et R30 : à confirmer.",
          "_etat/DECISIONS.md, entrée du 03/10 sur le 4H : « ± 0,067 » à remplacer par 0,066 (valeur corrigée le 05/10).",
          "Hygiène du dépôt : CLAUDE.md périmé, doublons à la racine, environ 1 Go de tables ASV redondantes, nom historique fig_S3_permdisp.png écrit par scripts/41."]:
    B(t)

# ------------------------------------------------------------------ méthode
H("Méthode de la vérification des références", 1)
P("Liste de références extraite de docs/manuscrit/Article.docx sur meso le 07/10/2026 (51 références). Pour les 48 DOI "
  "présents, interrogation de l'API Crossref (api.crossref.org/works) : dates published-print, published-online et issued, "
  "volume, page ou article-number. Comparaison automatique avec l'année, le volume et la pagination cités (tirets "
  "normalisés). Dates de publication CRAN lues dans les fichiers DESCRIPTION des paquets installés sur meso. Non vérifiés : "
  "listes d'auteurs, titres, rétractations, thèse [26], dépôt [51]. Table détaillée : docs/biblio/verif_crossref_2026-10-07.csv.")


# --- post-traitement accessibilité : numPr explicite sur les listes, langue fr-FR partout
def style_numid(name):
    pPr = doc.styles[name].element.pPr
    np_ = pPr.find(qn("w:numPr")) if pPr is not None else None
    return np_.find(qn("w:numId")).get(qn("w:val")) if np_ is not None else None
for p in doc.paragraphs:
    if p.style.name in ("List Bullet", "List Number"):
        nid = style_numid(p.style.name)
        if nid:
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr"); il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
            ni = OxmlElement("w:numId"); ni.set(qn("w:val"), nid); numPr.append(il); numPr.append(ni); pPr.append(numPr)
rpr_def = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
l = rpr_def.find(qn("w:lang"))
if l is None: l = OxmlElement("w:lang"); rpr_def.append(l)
l.set(qn("w:val"), "fr-FR")
st = doc.settings.element
tfl = st.find(qn("w:themeFontLang"))
if tfl is None: tfl = OxmlElement("w:themeFontLang"); st.append(tfl)
tfl.set(qn("w:val"), "fr-FR")

doc.save(OUT)
print("ok", OUT)
