"""49-d2_sejour_vivier.py — D2 tranchée par JF le 2026-10-06, option (ii).

L'article propose désormais la durée de séjour en vivier avant dissection comme mécanisme
candidat de l'effet de colonne, au lieu de laisser le mécanisme non identifié. La prémisse qui
départage — la position sur plaque n'a pas d'effet mesurable sur la composition — est une
observation non publiée du laboratoire (JF, 2026-10-06) ; elle est extérieure à ce jeu de
données, où R32 montre que colonne, ordre de traitement et durée de séjour ne sont pas
séparables (ρ de Spearman 0,548–0,950 sur 12 campagnes station × année).

Quatre modifications sur docs/manuscrit/Article.docx (version resserrée, référence unique
depuis le 2026-10-06) :
  1. Results, « Plate position » : « the delay between capture and dissection » ne correspond
     pas au protocole (poissons gardés vivants en vivier, dissection en 4–6 min) → remplacé par
     la durée de séjour en vivier.
  2. Results, paragraphe suivant : lecture du mécanisme ajoutée (témoin colonne/rang/bord +
     observation non publiée), avec mention explicite du statut d'hypothèse.
  3. Discussion, paragraphe des limites : même lecture, même réserve.
  4. Note grise « À reformuler après l'arbitrage D2 de JF » supprimée (devenue sans objet) ;
     une note grise de traçabilité de l'arbitrage la remplace, à retirer avant soumission.

python-docx n'est présent sur aucun interpréteur de meso : ce script s'exécute hors cluster,
comme les scripts 42, 44, 47 et 48 ; le .docx produit est redéposé dans le dépôt et contrôlé
sur meso par les scripts 40 et 30.
"""
import docx, hashlib, os
from docx.shared import RGBColor, Pt

M = os.environ.get("MS_DIR", "docs/manuscrit/")
OUT = os.environ.get("OUT_DIR", M)
F = "Article.docx"
MD5_IN = "1baaab97fde111200770a4160190e6d4"   # version resserrée promue en référence (commit a2752fd)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x88, 0x88, 0x88)


def para(doc, signature):
    """Le paragraphe unique contenant `signature`."""
    P = [p for p in doc.paragraphs if signature in p.text]
    assert len(P) == 1, (signature, len(P))
    return P[0]


def replace_across_runs(p, old, new):
    """Remplace `old` par `new` dans un paragraphe, même si la chaîne traverse plusieurs runs.
    La mise en forme retenue est celle du premier run touché (tous bleus ici)."""
    runs = list(p.runs)
    texts = [r.text for r in runs]
    full = "".join(texts)
    assert full.count(old) == 1, (old[:50], full.count(old))
    start = full.index(old)
    end = start + len(old)
    pos, done = 0, False
    for r, t in zip(runs, texts):
        s, e = pos, pos + len(t)
        pos = e
        if e <= start or s >= end:
            continue
        head = t[:start - s] if s < start else ""
        tail = t[end - s:] if e > end else ""
        r.text = (head + new + tail) if not done else (head + tail)
        done = True
    assert "".join(r.text for r in p.runs).count(new) == 1


assert hashlib.md5(open(M + F, "rb").read()).hexdigest() == MD5_IN, "md5 d'entrée inattendu"
d = docx.Document(M + F)

# ── 1. Results, « Plate position » : protocole ───────────────────────────────────────────────
p = para(d, "Three witnesses separate this effect")
replace_across_runs(
    p,
    "and the delay between capture and dissection are not separated by these data",
    "and the time each fish spent in the holding tank before dissection are not separated by "
    "these data",
)

# ── 2. Results : lecture du mécanisme ────────────────────────────────────────────────────────
p79 = para(d, "Within every fishing campaign, plate column followed the order")
replace_across_runs(
    p79,
    "and therefore from the time each fish spent in the holding tank before dissection.",
    "and therefore from the time each fish spent in the holding tank before dissection. Two "
    "observations make a physical plate artefact the less likely reading: the effect is specific "
    "to column, row reaching significance in only 7 and 6 of 24 strata and edge in none, and "
    "position within a plate has shown no measurable effect on 16S community composition in our "
    "laboratory (unpublished observations). We therefore read the column effect as a marker of "
    "processing order and propose the time spent in the holding tank before dissection as the "
    "candidate mechanism; this design cannot test that proposal, which remains a hypothesis.",
)

# ── 3. Discussion, paragraphe des limites ────────────────────────────────────────────────────
replace_across_runs(
    para(d, "Four limits bound these conclusions."),
    "is not identified, and because C. nasus was generally processed first it could contribute to",
    "is not identified by these data; position within a plate has no measurable effect on 16S "
    "composition in our laboratory (unpublished observations) and the effect was specific to "
    "column, so we read it as a marker of processing order and propose the time spent in the "
    "holding tank before dissection as the candidate mechanism, a hypothesis this design cannot "
    "test. Because C. nasus was generally processed first, processing order could contribute to",
)

# ── 4. Notes grises : l'ancienne devient sans objet, une note de traçabilité la remplace ─────
old_note = para(d, "À reformuler après l'arbitrage D2 de JF")
old_note._element.getparent().remove(old_note._element)

g = p79.add_run(
    " [Mécanisme proposé sur arbitrage de JF du 2026-10-06 (D2, option ii). La prémisse — la "
    "position sur plaque n'a pas d'effet mesurable sur la composition — est une observation non "
    "publiée du laboratoire, extérieure à ce jeu de données : R32 montre que colonne, ordre de "
    "traitement et durée de séjour ne sont pas séparables ici. À retirer avant soumission.]"
)
g.italic = True
g.font.size = Pt(9)
g.font.color.rgb = GREY

d.save(OUT + F)
print("écrit :", OUT + F, hashlib.md5(open(OUT + F, "rb").read()).hexdigest())
