"""50-retention_trois_runs.py — décision 5 de JF (2026-10-06) : les écarts de rétention selon la
profondeur sont rapportés comme l'étendue sur les trois runs de séquençage, et non plus comme les
valeurs du seul run durance1.

Source unique des valeurs : results/depth_agreement/retention_par_categorie.tsv (scripts/43), runs
durance1–3 (la ligne « tous » n'est pas utilisée). Valeurs de la station Confluence Buech-Meouge :
retention_par_station.tsv. Les chaînes insérées sont construites depuis ces tables et comparées aux
chaînes attendues (assertions) : aucun chiffre n'est retapé.

Modifications :
  Article.docx, §8.6 — trois phrases chiffrées, note grise « à trancher par JF » remplacée par une note
    de traçabilité.
  Supplementary_Data.docx — Note S2 (« Depth and sample selection », texte déplacé du §8.6 le 05/10),
    note et dernière colonne de la Table S6, légende de la Figure S1 (la figure elle-même trace déjà
    durance1 en ligne et l'étendue des trois runs en bande : seule la légende change).

Les remplacements se font à l'intérieur d'un run (les noms d'espèce sont des runs italiques
séparés). python-docx est absent de meso : exécution hors cluster, comme 42, 44, 47–49.
"""
import docx, hashlib, os
import pandas as pd
from docx.shared import RGBColor, Pt

M = os.environ.get("MS_DIR", "docs/manuscrit/")
OUT = os.environ.get("OUT_DIR", M)
R = os.environ.get("RES_DIR", "results/depth_agreement/")
MD5 = {"Article.docx": "ab68013b7ad97e913742ae85c4e69e07",           # commit 0bf81a7
       "Supplementary_Data.docx": "fb75d3e21e6cb3379e1cf4a654171367"}  # commit a2752fd
GREY = RGBColor(0x88, 0x88, 0x88)
RUNS = ["durance1", "durance2", "durance3"]

RT = pd.read_csv(R + "retention_par_categorie.tsv", sep="\t")
RT = RT[RT.run.isin(RUNS)].copy()
# Valeurs exactes depuis les effectifs : les colonnes du script 43 sont arrondies à 2 décimales, et
# les réarrondir à 1 décimale produit un double arrondi (32,654 -> 32,65 -> 32,6 au lieu de 32,7, passe 1,
# durance1, 2 000 lectures — seul cas de la grille ; aucune valeur citée dans l'article n'est concernée).
RT["retention_Hy"] = 100 * RT.retenus_Hy / RT.n_Hy
RT["retention_Pt"] = 100 * RT.retenus_Pt / RT.n_Pt
RT["ecart_Pt_moins_Hy"] = RT.retention_Pt - RT.retention_Hy
RS = pd.read_csv(R + "retention_par_station.tsv", sep="\t")
RS = RS[RS.run.isin(RUNS)]


def rng(df, col, **k):
    x = df
    for a, v in k.items():
        x = x[x[a] == v]
    assert len(x) == 3 and set(x.run) == set(RUNS), (k, len(x))
    return float(x[col].min()), float(x[col].max())


def fr(lohi, f):
    a, b = format(lohi[0], f), format(lohi[1], f)
    return a if a == b else f"{a}–{b}"


C = dict(passe=1, tissu="caudale", profondeur=3000)
h1 = fr(rng(RT, "retention_Hy", **C), ".0f"); p1 = fr(rng(RT, "retention_Pt", **C), ".0f")
h2 = fr(rng(RT, "retention_Hy", **{**C, "passe": 2}), ".0f")
g1 = fr(rng(RT, "ecart_Pt_moins_Hy", **C), ".0f"); g2 = fr(rng(RT, "ecart_Pt_moins_Hy", **{**C, "passe": 2}), ".0f")
c1 = fr(rng(RT, "ecart_Pt_moins_Hy", passe=1, tissu="caudale", profondeur=500), ".1f")
c2 = fr(rng(RT, "ecart_Pt_moins_Hy", passe=2, tissu="caudale", profondeur=500), ".1f")
m1 = fr(rng(RT, "ecart_Pt_moins_Hy", passe=1, tissu="midgut", profondeur=500), ".1f")
m2 = fr(rng(RT, "ecart_Pt_moins_Hy", passe=2, tissu="midgut", profondeur=500), ".1f")
SB = dict(station="Confluence Buech-Meouge", tissu="caudale", profondeur=3000)
s1 = fr(rng(RS, "ecart_Pt_moins_Hy", passe=1, **SB), ".0f"); s2 = fr(rng(RS, "ecart_Pt_moins_Hy", passe=2, **SB), ".0f")
g1d = fr(rng(RT, "ecart_Pt_moins_Hy", **C), ".1f"); g2d = fr(rng(RT, "ecart_Pt_moins_Hy", **{**C, "passe": 2}), ".1f")

# chaînes attendues, d'après R35 (RESULTATS.md) — l'assertion échoue si la table a changé
assert (h1, p1, h2, g1, g2) == ("50", "79–86", "64–67", "29–36", "15–20"), (h1, p1, h2, g1, g2)
assert (c1, c2, m1, m2) == ("2.6–3.8", "3.4–5.7", "5.7–14.0", "7.8–8.9"), (c1, c2, m1, m2)
assert (s1, s2, g1d, g2d) == ("20–48", "7–36", "29.0–36.4", "14.7–19.8"), (s1, s2, g1d, g2d)


def in_run(p, old, new):
    """Remplace `old` par `new` dans l'unique run qui le contient (mise en forme conservée)."""
    hits = [r for r in p.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (old[:60], len(hits))
    hits[0].text = hits[0].text.replace(old, new)


def para(doc, start):
    P = [p for p in doc.paragraphs if p.text.startswith(start)]
    assert len(P) == 1, (start, len(P))
    return P[0]


def grey_note(p, starts, text):
    """Remplace le texte du run gris qui commence par `starts`."""
    G = [r for r in p.runs if r.text.lstrip().startswith(starts)]
    assert len(G) == 1, (starts, len(G))
    G[0].text = text


TRACE = (" [Écarts de rétention : étendue sur les trois runs de séquençage, arbitrage de JF du 2026-10-06 "
         "(décision 5), au lieu des valeurs du seul run durance1. Source : scripts/43-retention_profondeur.py, "
         "results/depth_agreement/retention_par_categorie.tsv ; insertion par scripts/50-retention_trois_runs.py. "
         "À retirer avant soumission.]")

# ── Article, §8.6 ────────────────────────────────────────────────────────────────────────────
for f in MD5:
    assert hashlib.md5(open(M + f, "rb").read()).hexdigest() == MD5[f], f"md5 d'entrée inattendu : {f}"

a = docx.Document(M + "Article.docx")
p = para(a, "Because the threshold also decides")
in_run(p, "in the caudal fin, in sequencing run durance1, 50 % of intermediate hybrids",
       f"in the caudal fin, across the three sequencing runs, {h1} % of intermediate hybrids")
in_run(p, " passed the threshold against 86 % of ", f" passed the threshold against {p1} % of ")
in_run(p, "and 67 % of all 42 hybrids (pass 2), a gap of 36 and 20 points respectively.",
       f"and {h2} % of all 42 hybrids (pass 2), a gap of {g1} and {g2} points respectively.")
in_run(p, "falls to 3.8 points (pass 1) and 5.7 points (pass 2)", f"falls to {c1} points (pass 1) and {c2} points (pass 2)")
in_run(p, "the midgut gap (14.0 and 8.9 points)", f"the midgut gap ({m1} and {m2} points)")
grey_note(p, "[Écarts de rétention : valeurs du run durance1", TRACE)
a.save(OUT + "Article.docx")

# ── Supplément ───────────────────────────────────────────────────────────────────────────────
s = docx.Document(M + "Supplementary_Data.docx")
p = para(s, "Depth and sample selection.")
in_run(p, "in the caudal fin, in sequencing run durance1, 50 % of intermediate hybrids",
       f"in the caudal fin, across the three sequencing runs, {h1} % of intermediate hybrids")
in_run(p, " passed the threshold against 86 % of ", f" passed the threshold against {p1} % of ")
in_run(p, "and 67 % of all 42 hybrids (pass 2), a gap of 36 and 20 points respectively, reaching 48 and 36 points at one",
       f"and {h2} % of all 42 hybrids (pass 2), a gap of {g1} and {g2} points respectively, and of {s1} and {s2} points at one")
in_run(p, "falls to 3.8 points (pass 1) and 5.7 points (pass 2)", f"falls to {c1} points (pass 1) and {c2} points (pass 2)")
in_run(p, "the midgut gap (14.0 and 8.9 points)", f"the midgut gap ({m1} and {m2} points)")

p = para(s, "Pearson correlation with the 3,000-read reference")
in_run(p, "in pass 1 / pass 2; at 3,000 reads it is 36.4 / 19.8 points.",
       f"in pass 1 / pass 2, given as the range across the three sequencing runs; at 3,000 reads it is {g1d} / {g2d} points.")

p = para(s, "Figure S1.")
in_run(p, "Lines: sequencing run durance1 (values quoted in Materials and Methods §8.6); shaded bands: range across the three runs.",
       "Lines: sequencing run durance1; shaded bands: range across the three runs (quoted in Materials and Methods §8.6).")

T = [t for t in s.tables if t.rows[0].cells[-1].text.startswith("Retention gap Pt−Hy, caudal fin")]
assert len(T) == 1, len(T)
t = T[0]
hdr = t.rows[0].cells[-1].paragraphs[0]
in_run(hdr, "pass 1 / pass 2", "pass 1 / pass 2, range across runs")
for row in t.rows[1:]:
    depth = int(float(row.cells[0].text))
    v1 = fr(rng(RT, "ecart_Pt_moins_Hy", passe=1, tissu="caudale", profondeur=depth), ".1f")
    v2 = fr(rng(RT, "ecart_Pt_moins_Hy", passe=2, tissu="caudale", profondeur=depth), ".1f")
    cp = row.cells[-1].paragraphs[0]
    old = cp.text
    d1 = fr((float(RT[(RT.run == "durance1") & (RT.passe == 1) & (RT.tissu == "caudale") & (RT.profondeur == depth)].ecart_Pt_moins_Hy.iloc[0]),) * 2, ".1f")
    assert old.startswith(d1 + " / "), (depth, old, d1)        # l'ancienne cellule était bien durance1
    in_run(cp, old, f"{v1} / {v2}")
    print(f"Table S6, {depth} lectures : {old!r} -> {v1} / {v2}")
s.save(OUT + "Supplementary_Data.docx")

for f in MD5:
    print("écrit :", OUT + f, hashlib.md5(open(OUT + f, "rb").read()).hexdigest())
