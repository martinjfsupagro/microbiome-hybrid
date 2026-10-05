"""48-correction_camper_SI.py — §8.11 : « 0.067 » -> « 0.066 », d'après les Supporting Information de Camper et al. 2024
(Table S.4.1, lézard Aspidoscelis, version Jaccard : axe parental 0,4183 à 1 000 lectures, 0,4847 à 10 000, soit 0,0664),
extraites dans docs/biblio/camper2024_SI_coeur_profondeur.tsv. Appliqué aux deux versions de l'article (gardes md5)."""
import docx, hashlib, os, copy
from docx.shared import RGBColor, Pt
from docx.text.run import Run
M = os.environ.get("MS_DIR", "docs/manuscrit/"); OUT = os.environ.get("OUT_DIR", M)
IN = {"Article.docx": "553f9904886b8856517ca5567f68089f", "Article_resserre.docx": "b77ff31ef8d85e7ed10987bb293a5b72"}
OLD = "varies by at most 0.067 between 1,000 and 10,000 reads [14]"
NEW = "varies by at most 0.066 between 1,000 and 10,000 reads [14]"
NOTE = (" [0,066 : Camper et al. 2024, Supporting Information, Table S.4.1 (lézard Aspidoscelis, version Jaccard, axe parental "
        "0,4183 → 0,4847) ; 0,045 en version Bray-Curtis (Table S.4.2). L'ancienne valeur 0,067 était mal arrondie. Exclusion de "
        "ρ ≥ 0,8 vérifiée : Intersection = 0 à ρ = 0,8 dans les Tables S.2.1 et S.2.2 (0,015 et 0,010 à ρ = 0,7). Vérifié le 2026-10-05.]")
for f, h in IN.items():
    assert hashlib.md5(open(M + f, "rb").read()).hexdigest() == h, f
    d = docx.Document(M + f)
    P = [p for p in d.paragraphs if OLD in p.text]; assert len(P) == 1, f; p = P[0]
    R = [r for r in p.runs if OLD in r.text]; assert len(R) == 1, f; r = R[0]
    a, b = r.text.split(OLD, 1); r.text = a
    n1 = copy.deepcopy(r._r); r._r.addnext(n1); n2 = copy.deepcopy(r._r); n1.addnext(n2)
    R1, R2 = Run(n1, p), Run(n2, p); R1.text = NEW; R1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); R2.text = b
    g = p.add_run(NOTE); g.italic = True; g.font.size = Pt(9); g.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    d.save(OUT + f); print(f, "ok")
