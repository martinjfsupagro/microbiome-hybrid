"""52-reponses_andre_1006.py — réponses d'André du 2026-10-06 (docs/reponses_andre_2026-10-06.md) portées au manuscrit.

Article.docx
  §1 (Methods)      euthanasie : raison donnée par André ; autorités des deux autorisations et organismes
                    du Suran (Fédération de l'Ain) et de l'Ardèche (ONEMA 07) ; note grise remplacée.
  §2 (Methods)      conservation du tube éthanol : glace sur le terrain, −80 °C au laboratoire ; note grise remplacée.
  Ethics            même contenu que §1, méthode d'euthanasie explicite ; note grise remplacée.
  Funding           référence au projet FACIES retirée (André : « on enleve la reference au projet facies »),
                    ainsi que « the contract that also funded Guivier et al. [31] », qui reposait sur ce même
                    rattachement ; note grise remplacée.
  Results (alpha)   note grise mise à jour (code X ; test R31 fait) — texte du paragraphe inchangé.
Supplementary_Data.docx
  Table S2          légende : codes de sexe définis ; note grise mise à jour.

Tout texte nouveau est inséré dans un run bleu (1F4E79), convention du manuscrit pour le texte ajouté ;
le run d'origine est découpé si nécessaire (comme scripts/48). Gardes md5 en entrée.
python-docx absent de meso : exécution hors cluster.
"""
import docx, hashlib, os, copy
from docx.shared import RGBColor
from docx.text.run import Run

M = os.environ.get("MS_DIR", "docs/manuscrit/"); OUT = os.environ.get("OUT_DIR", M)
MD5 = {"Article.docx": "cbf412a73c633ac03870d6b9c60193f2",            # commit f7b41ba
       "Supplementary_Data.docx": "20f30aca78ce57a2adcbf100329fb1f7"}
BLUE = RGBColor(0x1F, 0x4E, 0x79)


def para(doc, start):
    P = [p for p in doc.paragraphs if p.text.startswith(start)]
    assert len(P) == 1, (start, len(P)); return P[0]


def blue_replace(p, old, new):
    """Remplace `old` (contenu dans un seul run) par `new` placé dans un run bleu ; le reste du run garde sa forme."""
    R = [r for r in p.runs if old in r.text]
    assert len(R) == 1 and R[0].text.count(old) == 1, (old[:60], len(R))
    r = R[0]
    a, b = r.text.split(old, 1)
    r.text = a
    n1 = copy.deepcopy(r._r); r._r.addnext(n1); n2 = copy.deepcopy(r._r); n1.addnext(n2)
    R1, R2 = Run(n1, p), Run(n2, p)
    R1.text = new; R1.font.color.rgb = BLUE
    R2.text = b


def grey(p, starts, text):
    G = [r for r in p.runs if r.text.lstrip().startswith(starts)]
    assert len(G) == 1, (starts, len(G)); G[0].text = text


for f, h in MD5.items():
    assert hashlib.md5(open(M + f, "rb").read()).hexdigest() == h, f"md5 d'entrée inattendu : {f}"

AUTH = ("under authorisations n° 2014-156-0001 and n° 2015-1426DDT605, approved by ONEMA and the Directions "
        "départementales des territoires of Alpes-de-Haute-Provence, Hautes-Alpes and Vaucluse; fishing in the "
        "Suran was conducted with the Fédération de l'Ain pour la pêche et la protection des milieux aquatiques, "
        "and in the Ardèche with the Ardèche departmental service of ONEMA.")
OUVERT = ("Reste ouvert : aucun numéro d'autorisation n'est donné pour le Suran ni pour l'Ardèche ; place de Benjamin "
          "Hérodet (Fédération de l'Ain) à préciser — déclaration d'éthique, remerciements ou auteurs.")

a = docx.Document(M + "Article.docx")

# ── §1 ──────────────────────────────────────────────────────────────────────────────────────
p = para(a, "Specimens were captured by electrofishing")
blue_replace(p, "each fish was euthanised by cervical dislocation and immediately dissected",
             "each fish was euthanised by cervical dislocation, the method retained so as not to alter gene "
             "expression in the tissues shared with transcriptome analyses (§2), and immediately dissected")
blue_replace(p, "and under national protocol authorisations n° 2014-156-0001 and n° 2015-1426DDT605.", "and " + AUTH)
grey(para(a, "[À confirmer par André — questions de suivi envoyées le 2026-10-03. (1)"),
     "[À confirmer par André",
     "[Autorités, euthanasie et raison de la méthode : André, 2026-10-06 (docs/reponses_andre_2026-10-06.md). "
     + OUVERT + " La séquence pêche → vivier → dissection, les durées et l'ordre de traitement viennent d'André "
     "(2026-10-03 ; hotus en priorité : 2026-09-25).]")

# ── §2 ──────────────────────────────────────────────────────────────────────────────────────
p = para(a, "Four tissues were sampled per fish")
blue_replace(p, "one half was stored in 95% ethanol at −80 °C for microbiota analysis, the other",
             "one half was placed in 95% ethanol for microbiota analysis, kept on ice in the field and stored at "
             "−80 °C in the laboratory; the other")
grey(para(a, "[À confirmer par André — question de suivi envoyée le 2026-10-03 : température"),
     "[À confirmer par André",
     "[Conservation du tube éthanol (glace sur le terrain, −80 °C au laboratoire) : André, 2026-10-06. Tissu 01 = "
     "lobe de nageoire caudale, ordre de dissection et partage de chaque tissu en deux tubes : André, 2026-10-03. "
     "À retirer avant soumission.]")

# ── Ethics ──────────────────────────────────────────────────────────────────────────────────
p = para(a, "Fish were captured by electrofishing and euthanised before dissection")
blue_replace(p, "Fish were captured by electrofishing and euthanised before dissection under authorisations n° "
                "2014-156-0001 and n° 2015-1426DDT605, following international guidelines of animal care.",
             "Fish were captured by electrofishing and euthanised by cervical dislocation immediately before "
             "dissection, following international guidelines of animal care, " + AUTH)
grey(p, "[À compléter (questions de suivi à André, 2026-10-03)",
     " [Autorités et euthanasie : André, 2026-10-06. " + OUVERT + " Consentement : sans objet.]")

# ── Funding ─────────────────────────────────────────────────────────────────────────────────
p = para(a, "This work was funded by Électricité de France")
blue_replace(p, "financing the thesis of Arnaud Ungaro, the contract that also funded Guivier et al. [31] within the "
                "FACIES project, with the support", "financing the thesis of Arnaud Ungaro, with the support")
grey(para(a, "[Financement : référence du contrat"), "[Financement",
     "[Financement : contrat « EDF-CNRS AGDI 428481 » (André, 2026-10-03). Référence au projet FACIES retirée sur "
     "décision d'André (2026-10-06 : rattachement non établi) ; la mention « the contract that also funded Guivier "
     "et al. [31] », qui reposait sur ce même rattachement, est retirée avec elle — à valider par JF. Section "
     "Acknowledgements : à rédiger (André) ; y nommer Benjamin Hérodet si c'est la place retenue.]")

# ── Results, alpha : note grise seulement ───────────────────────────────────────────────────
grey(para(a, "Hybrids were transgressive in none of the 32 combinations"), "[Non testé : C. nasus",
     " [Ordre de traitement testé le 2026-10-06 (plan docs/plan_R31_ordre_2026-10-06.md, f3916c1 ; scripts/51) : "
     "dominance de C. nasus non attribuable à l'ordre ; intégration au texte en attente de JF. Code de sexe « X » : "
     "individus non sexés (André, 2026-10-06) ; le contraste femelles/mâles de R27 (F 33, M 78 ; X et vides exclus) "
     "peut être rapporté — décision de JF attendue.]")
a.save(OUT + "Article.docx")

# ── Supplément, Table S2 ────────────────────────────────────────────────────────────────────
s = docx.Document(M + "Supplementary_Data.docx")
p = para(s, "Genotypic category of each of the 180 individuals")
blue_replace(p, "TL: total length.",
             "TL: total length. Sex: F, female; M, male; X and –, sex not determined (juveniles, or sex not assigned).")
grey(para(s, "[À confirmer : signification du code de sexe"), "[À confirmer",
     "[Codes de sexe : X (52 individus) et – (17) désignent des individus non sexés ; selon André (2026-10-06), "
     "l'une des deux catégories regroupe des juvéniles, l'autre des individus dont le sexe n'a pas été attribué — "
     "laquelle est laquelle : [À CONFIRMER]. Table insérée le 2026-10-03 sur décision de JF (question 2.9) ; les "
     "tables suivantes sont renumérotées S3 à S10.]")
s.save(OUT + "Supplementary_Data.docx")

for f in MD5:
    print("écrit :", OUT + f, hashlib.md5(open(OUT + f, "rb").read()).hexdigest())
