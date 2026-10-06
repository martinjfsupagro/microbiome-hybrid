#!/usr/bin/env python3
"""74-maj_image_fig_S3.py — remplace, dans la version de travail du supplément, l'image de la Figure S3 par le rendu
à phylums en italique (titres et légende ; consigne de la revue : tous les rangs taxonomiques en italique).

Usage : 74-maj_image_fig_S3.py <Supplementary_Data.docx> <fig_S3_phyla.png nouveau> <md5 attendu de l'ancienne image> <sortie.docx>
Garde-fous : l'image remplacée est celle qui suit le titre de la Figure S3, référencée une seule fois, d'empreinte attendue,
et de mêmes dimensions en pixels que la nouvelle (le cadre du docx n'est pas modifié).
"""
import sys, os, io, re, hashlib
import docx
from lxml import etree
from PIL import Image
from docx.opc.constants import RELATIONSHIP_TYPE as RT
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import soumission_commun as SC

SUP, PNG, MD5_ANC, OUT = sys.argv[1:5]
S = docx.Document(SUP)
body, EL = SC.elements_supplement(S)
k0 = next(e[0] for e in EL if e[2] == "F3")
rid = re.search(r'r:embed="(rId\d+)"', etree.tostring(body[k0]).decode()).group(1)
assert sum(f'r:embed="{rid}"' in etree.tostring(b).decode() for b in body) == 1
part = S.part.rels[rid].target_part
assert hashlib.md5(part.blob).hexdigest() == MD5_ANC, hashlib.md5(part.blob).hexdigest()
neuf = open(PNG, "rb").read()
assert Image.open(io.BytesIO(part.blob)).size == Image.open(io.BytesIO(neuf)).size
part._blob = neuf
S.save(OUT)
v = docx.Document(OUT)
assert hashlib.md5(v.part.rels[rid].target_part.blob).hexdigest() == hashlib.md5(neuf).hexdigest()
assert [p.text for p in v.paragraphs] == [p.text for p in docx.Document(SUP).paragraphs]
print(OUT, "| md5", hashlib.md5(open(OUT, "rb").read()).hexdigest(), "| image", rid, MD5_ANC, "->", hashlib.md5(neuf).hexdigest())
